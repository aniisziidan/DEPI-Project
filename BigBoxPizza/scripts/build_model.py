"""Build BigBoxPizza_US_Landed_Cost_Model.xlsx (formula-driven) with openpyxl.

Run:  python3 build_model.py [output_path]
Inputs: inputs.py (all assumptions), loading_model.py (static DP/optimiser values), sources_data.py,
        registers.py.  All key financial results are Excel formulas; only (a) inputs and (b) a few
        script-calculated geometry values (orange cells, documented) are constants.
"""
import os
import sys

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment

import engine as E
import inputs as I
import loading_model as LM
import registers as R
import sources_data as SD
from xlkit import *  # noqa: F401,F403
import build_views as V

HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.path.dirname(HERE)

SH = dict(
    dash="Executive Dashboard", readme="Read Me & Instructions", sku="Product & SKU Master", quote="Supplier Quotation",
    spec="Technical Specifications", fx="Exchange Rates", cont="Container Specifications", geo="Loading Geometry",
    ship="Shipment Configuration", orig="Origin Costs", ocean="Ocean Freight", usc="U.S. Port & Customs",
    dest="Destination Delivery", tariff="Tariffs & Customs Value", landed="Landed Cost by SKU", price="Pricing & Margin",
    prof="Container Profitability", sens="Sensitivity Analysis", scen="Scenario Comparison", cash="Cash Flow & Working Capital",
    comp="Competitor Benchmark", risk="Risk Register", assum="Assumptions Register", src="Source Register",
    rfq="Quote Requests & Missing Data", eng="Calc Engine", chk="Model Checks",
)
ORDER = ["dash", "readme", "sku", "quote", "spec", "fx", "cont", "geo", "ship", "orig", "ocean", "usc", "dest", "tariff",
         "landed", "price", "prof", "sens", "scen", "cash", "comp", "risk", "assum", "src", "rfq", "eng", "chk"]
SKU_CODES = [k["code"] for k in I.SKUS]
CT_KEYS = ["20GP", "40GP", "40HC"]
CT_IDX = {"20GP": 1, "40GP": 2, "40HC": 3, "LCL": 4}
PORTS = [("NYNJ", "New York / New Jersey"), ("BAL", "Baltimore"), ("ORF", "Norfolk (Port of Virginia)"), ("SAV", "Savannah")]
PORT_IDX = {p[0]: i + 1 for i, p in enumerate(PORTS)}


# =============================================================================== input sheets
def build_quote(book):
    ws = book.ws(SH["quote"])
    title(ws, "Supplier Quotation – NFC-OFFER-2026-0120 (as printed) with arithmetic checks",
          "Source S001: BigBoxPizza/evidence/NFC-OFFER-2026-0120_original.pdf (image PDF, transcribed). Blue = input as printed; grey = formula.")
    set_widths(ws, {"A": 26, "B": 30, "C": 11, "D": 9, "E": 9, "F": 9, "G": 10, "H": 12, "I": 15, "J": 15, "K": 11, "L": 12, "M": 13, "N": 13, "O": 14, "P": 14})
    r = section(ws, 4, "1. Quotation header (verbatim)", 16)
    hdr = [("Quotation number", I.QUOTE["number"]), ("Quotation date", "2026-09-23"), ("Valid until (printed)", "2026-09-30"),
           ("Currency", "EGP"), ("Tax rate printed", 0.14), ("Payment terms printed", "advance_50"),
           ("Incoterm printed", "NONE STATED (EXW is only the brief's working assumption)"), ("Shipping printed", "— (dash)"),
           ("Client printed", "ا/ انيس شركة توريدات – Alexandria – code C330 – tax ID blank"),
           ("Bank printed", "FAP MISR BANK, SWIFT NBADEGCAXXX, IBAN EG59 0019 0021 0000 0000 7554 9000 3 (mod-97 valid)"),
           ("Signatory", "Script 'Abdelrahman Darwish' over printed name 'Mahmoud Bakr'"), ("Supplier address", "199 6th Ext. Industrial Zone, 6 October City, Giza – Tax ID 552-782-209")]
    import datetime as dt
    for lab, val in hdr:
        put(ws, r, 1, lab, bold=True)
        if lab in ("Quotation date", "Valid until (printed)"):
            c = inp_cell(ws, r, 2, dt.date.fromisoformat(val), "yyyy-mm-dd")
            if lab.startswith("Valid"):
                book.name("q_valid_until", ws.title, f"B{r}")
            else:
                book.name("q_date", ws.title, f"B{r}")
        elif lab == "Tax rate printed":
            inp_cell(ws, r, 2, val, FMT_PCT)
            book.name("q_tax_rate", ws.title, f"B{r}")
        else:
            c = inp_cell(ws, r, 2, val)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        r += 1
    put(ws, r, 1, "Validity status at valuation date", bold=True)
    calc_cell(ws, r, 2, '=IF(val_date>q_valid_until,"EXPIRED "&TEXT(val_date-q_valid_until,"0")&" days before valuation date – prices are historical reference only","Valid")')
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    book.name("q_status", ws.title, f"B{r}")
    r += 2
    r = section(ws, r, "2. Line items (as printed) and arithmetic checks", 16)
    labels = ["Code", "Description (verbatim Arabic)", "Spec", "L mm", "W mm", "H mm", "Qty", "Unit EGP", "Line total EGP (printed)",
              "Line total = Qty × Unit", "Difference", "USD per box (active FX)", "USD line value", "EGP per 1,000", "EGP per m² plan area", "Unit EGP incl. 14% tax"]
    r = header(ws, r, labels)
    first = r
    for k in I.SKUS:
        put(ws, r, 1, k["code"], bold=True)
        put(ws, r, 2, k["desc_ar"])
        put(ws, r, 3, "SINGLE-E")
        for j, key in enumerate(["L", "W", "H"]):
            inp_cell(ws, r, 4 + j, k[key], FMT_INT)
            book.name(f"q_{key}_{k['code']}", ws.title, f"{col(4 + j)}{r}")
        inp_cell(ws, r, 7, k["qty"], FMT_INT)
        inp_cell(ws, r, 8, k["price_egp"], "#,##0.000")
        inp_cell(ws, r, 9, k["line_total"], FMT_NUM)
        calc_cell(ws, r, 10, f"=G{r}*H{r}", FMT_NUM)
        calc_cell(ws, r, 11, f"=J{r}-I{r}", FMT_NUM)
        calc_cell(ws, r, 12, f"=H{r}/fx_active", FMT_USD4)
        calc_cell(ws, r, 13, f"=J{r}/fx_active", FMT_USD)
        calc_cell(ws, r, 14, f"=H{r}*1000", FMT_NUM)
        calc_cell(ws, r, 15, f"=H{r}/(D{r}*E{r}/1000000)", FMT_NUM)
        calc_cell(ws, r, 16, f"=H{r}*(1+q_tax_rate)", "#,##0.000")
        book.name(f"q_qty_{k['code']}", ws.title, f"G{r}")
        book.name(f"q_price_{k['code']}", ws.title, f"H{r}")
        book.name(f"q_total_{k['code']}", ws.title, f"I{r}")
        r += 1
    last = r - 1
    book.name("q_diff", ws.title, f"K{first}:K{last}")
    put(ws, r, 1, "Totals", bold=True)
    calc_cell(ws, r, 7, f"=SUM(G{first}:G{last})", FMT_INT, bold=True)
    calc_cell(ws, r, 10, f"=SUM(J{first}:J{last})", FMT_NUM, bold=True)
    calc_cell(ws, r, 13, f"=SUM(M{first}:M{last})", FMT_USD, bold=True)
    book.name("q_qty_total", ws.title, f"G{r}")
    book.name("q_sub_calc", ws.title, f"J{r}")
    r += 2
    checks = [("Subtotal (printed)", 463300.0, "=q_sub_calc", "Subtotal"),
              ("Tax 14% (printed)", 64862.0, "=q_sub_calc*q_tax_rate", "Tax"),
              ("Total (printed)", 528162.0, "=q_sub_calc*(1+q_tax_rate)", "Total")]
    r = header(ws, r, ["Check", "Printed EGP", "Recomputed EGP", "Difference", "Status"])
    for lab, printed, f, nm in checks:
        put(ws, r, 1, lab)
        inp_cell(ws, r, 2, printed, FMT_NUM)
        calc_cell(ws, r, 3, f, FMT_NUM)
        calc_cell(ws, r, 4, f"=C{r}-B{r}", FMT_NUM)
        calc_cell(ws, r, 5, f'=IF(ABS(D{r})<0.005,"OK","FAIL")')
        book.name(f"q_chk_{nm.lower()}", ws.title, f"E{r}")
        if nm == "Subtotal":
            book.name("q_subtotal", ws.title, f"B{r}")
        r += 1
    put(ws, r, 1, "Total quantity = 60,000")
    calc_cell(ws, r, 3, "=q_qty_total", FMT_INT)
    calc_cell(ws, r, 5, '=IF(q_qty_total=60000,"OK","FAIL")')
    book.name("q_chk_qty", ws.title, f"E{r}")
    r += 2
    r = section(ws, r, "3. Price sanity statistics (formula)", 16)
    stats = [
        ("Weighted average quoted price, EGP/box (excl. tax)", "=q_sub_calc/q_qty_total", "#,##0.0000"),
        ("Weighted average, USD/box at active FX", "=q_sub_calc/q_qty_total/fx_active", FMT_USD4),
        ("Weighted average incl. 14% tax, EGP/box", "=q_sub_calc*(1+q_tax_rate)/q_qty_total", "#,##0.0000"),
        ("Total quoted product value, USD (excl. tax, active FX)", "=q_sub_calc/fx_active", FMT_USD),
        ("Quoted value per 1,000 boxes, EGP", "=q_sub_calc/q_qty_total*1000", FMT_NUM),
        ("3464 (355 mm) vs 3463 (405 mm) unit price difference", "=q_price_3464/q_price_3463-1", FMT_PCT),
        ("3460 (305 mm) vs 3463 (405 mm) unit price difference", "=q_price_3460/q_price_3463-1", FMT_PCT),
        ("ANOMALY FLAG: smaller 3464 priced above larger 3463?", '=IF(q_price_3464>q_price_3463,"WARN – 355 mm box costs more than 405 mm box; ask NFC to explain (spec/print/error)","OK")', None),
    ]
    for lab, f, fm in stats:
        put(ws, r, 1, lab, wrap=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        calc_cell(ws, r, 6, f, fm)
        if "ANOMALY" in lab:
            ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=14)
            book.name("q_anomaly", ws.title, f"F{r}")
        r += 1
    r += 1
    r = section(ws, r, "4. Supplier commercial & tax inputs (scenario-dependent)", 10)
    items = list(I.SUPPLIER.items())
    rows, r = input_table(book, ws, r, "sup", items, "SUP")
    r += 1
    r = section(ws, r, "5. Quotation clarification checklist for NFC (send with RFQ 1)", 10)
    qs = R.NFC_CLARIFICATIONS
    for i, t in enumerate(qs, 1):
        put(ws, r, 1, f"Q{i:02d}")
        put(ws, r, 2, t, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=12)
        ws.row_dimensions[r].height = 28
        r += 1


def build_fx(book):
    ws = book.ws(SH["fx"])
    title(ws, "Exchange Rates – dated, editable inputs", "EGP per USD. BASE = Central Bank of Egypt buy rate 15-Oct-2026 (S063). Effective rate = CBE × (1 − bank spread).")
    set_widths(ws, {"A": 10, "B": 44, "C": 12, "D": 12, "E": 12, "F": 12, "G": 16, "H": 18, "I": 18, "J": 60})
    r = section(ws, 4, "1. FX inputs", 10)
    rows, r = input_table(book, ws, r, "fx", list(I.FX.items()), "FX")
    r += 1
    r = section(ws, r, "2. Effective EGP per USD by scenario", 10)
    r = header(ws, r, ["", "Measure", "", "LOW", "BASE", "STRESS", "ACTIVE"])
    put(ws, r, 2, "Effective EGP per USD = CBE × (1 − spread)")
    for j in range(3):
        calc_cell(ws, r, 4 + j, f"=INDEX(in_fx_cbe,1,{j + 1})*(1-INDEX(in_fx_spread,1,{j + 1}))", FMT_NUM4)
    calc_cell(ws, r, 7, f"=INDEX(D{r}:F{r},1,sel_scen)", FMT_NUM4, bold=True)
    book.name("fx_eff", ws.title, f"D{r}:F{r}")
    book.name("fx_active", ws.title, f"G{r}")
    r += 2
    r = section(ws, r, "3. Dated observations (evidence)", 10)
    r = header(ws, r, ["Date", "Observation", "EGP/USD", "", "", "", "", "Source", "Type", "Note"])
    obs = [("2026-02-22", "Pre-conflict level (press)", 47.64, "S064", "Press", ""),
           ("2026-03 (late)", "Dollar breaks 53.50; NBE quote 54.6", 54.60, "S064", "Press / bank quote", "Used as LOW (favourable to USD buyer)"),
           ("2026-05-14", "CBE ~52.83 buy / 52.97 sell", 52.83, "S064", "Press", ""),
           ("2026-07-16", "Pound firms below 51 (~50.6)", 50.60, "S064", "Press", ""),
           ("2026-10-11", "Aggregator 52.53 buy / 52.66 sell", 52.53, "S063", "Aggregator", ""),
           ("2026-10-14", "CBE homepage 53.0544 buy / 53.1544 sell", 53.0544, "S063", "Official (via search extract)", ""),
           ("2026-10-15", "CBE 53.04 buy / 53.18 sell (press report)", 53.04, "S063", "Official via press", "BASE"),
           ("2027-06 (forecast)", "Reuters poll: 49 by end-June 2027", 49.00, "S065", "Forecast", "Used as STRESS (EGP appreciation)")]
    for d, o, v, s, t, n in obs:
        put(ws, r, 1, d)
        put(ws, r, 2, o)
        static_cell(ws, r, 3, v, FMT_NUM4)
        put(ws, r, 8, s)
        put(ws, r, 9, t)
        put(ws, r, 10, n)
        r += 1
    put(ws, r + 1, 2, "FX date used for BASE", bold=True)
    inp_cell(ws, r + 1, 3, "2026-10-15")
    book.name("fx_date", ws.title, f"C{r + 1}")


def build_containers(book):
    ws = book.ws(SH["cont"])
    title(ws, "Container Specifications", "Representative dry-container data (DSV, S041; cross-check Maersk S042). US road cargo limits (S043, S096). Verify the actual unit's CSC plate.")
    set_widths(ws, {"A": 8, "B": 20, "C": 10, "D": 10, "E": 10, "F": 10, "G": 10, "H": 10, "I": 11, "J": 9, "K": 7, "L": 13, "M": 13, "N": 50, "O": 18})
    r = section(ws, 4, "1. Equipment (mm, kg, m³)", 15)
    r = header(ws, r, ["Key", "Container", "Inside L", "Inside W", "Inside H", "Door W", "Door H", "Tare kg", "Max payload kg", "Cube m³", "TEU",
                       "US road cargo limit kg", "Road limit lb", "Road-limit note", "Sources"])
    first = r
    for c in I.CONTAINERS:
        put(ws, r, 1, c["key"], bold=True)
        put(ws, r, 2, c["name"])
        for j, k in enumerate(["L", "W", "H", "door_W", "door_H", "tare", "payload"]):
            inp_cell(ws, r, 3 + j, c[k], FMT_INT)
        inp_cell(ws, r, 10, c["cube"], "0.0")
        inp_cell(ws, r, 11, c["teu"], "0")
        inp_cell(ws, r, 12, c["road_kg"], FMT_INT)
        calc_cell(ws, r, 13, f"=L{r}/0.45359237", FMT_INT)
        put(ws, r, 14, c["road_note"], wrap=True)
        put(ws, r, 15, I.CONTAINER_SOURCES)
        r += 1
    put(ws, r, 1, "LCL", bold=True)
    put(ws, r, 2, "LCL (consolidated)")
    inp_cell(ws, r, 11, 0, "0")
    put(ws, r, 14, "No container; TEU-based charges not applied")
    last = r - 1
    for nm, c in [("ct_L", "C"), ("ct_W", "D"), ("ct_H", "E"), ("ct_doorH", "G"), ("ct_road", "L"), ("ct_cube", "J")]:
        book.name(nm, ws.title, f"{c}{first}:{c}{last}")
    book.name("ct_teu", ws.title, f"K{first}:K{r}")
    r += 2
    r = section(ws, r, "2. Usable envelope and loading-efficiency inputs", 10)
    rows, r = input_table(book, ws, r, "use", list(I.USABLE.items()), "USE")
    r += 1
    notes = ["Usable length = inside length − deduction; usable width = inside width − deduction; usable height = MIN(inside height − top clearance, door height).",
             "Ocean payload (25–28.6 t) is never binding: the US road limit (~17.2 t 20', ~20.0 t 40'/HC on standard chassis) binds first.",
             "Loading efficiency reserves capacity for hand-stacked loose bundles (dimensional variation, dunnage, airbags). Required bundles are divided by efficiency before the geometric fit."]
    for t in notes:
        put(ws, r, 1, "Note", bold=True)
        put(ws, r, 2, t, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=14)
        ws.row_dimensions[r].height = 30
        r += 1


def build_geometry(book, static):
    ws = book.ws(SH["geo"])
    title(ws, "Loading Geometry – PROVISIONAL physical model (formula-driven)",
          "Blank sizes are a calibrated die-line ESTIMATE (2L + k·H by W + 2H + allowance), not NFC data. Replace with NFC die-lines/bundle data (RFQ 1). Orange = value written by scripts/loading_model.py.")
    widths = {c: 9 for c in "ABCDEFGHIJKLMNOPQRSTUVWXYZ"}
    widths.update({"A": 8, "B": 9, "AC": 18, "AH": 22, "AG": 10})
    set_widths(ws, widths)
    # ---------------- Block A: blanks & bundles
    r = section(ws, 4, "A. Flat blank and bundle per SKU and scenario", 24)
    labels = ["SKU", "Scenario", "Scen#", "L mm", "W mm", "H mm", "k (×H)", "W allow", "Blank L mm", "Blank W mm", "Bounding m²", "Net factor",
              "g/m²", "Blank kg", "Wrap mm", "Bundle a mm", "Bundle b mm", "Blanks/bundle", "Pitch mm", "Bundle h mm", "Bundle kg", "Bundle m³", "m³ per box", "kg per box"]
    r = header(ws, r, labels)
    geoA = {}
    for k in I.SKUS:
        code = k["code"]
        r_first = r
        for si, sn in enumerate(I.SCENARIOS, 1):
            put(ws, r, 1, code, bold=True)
            put(ws, r, 2, sn)
            put(ws, r, 3, si)
            calc_cell(ws, r, 4, f"=q_L_{code}", FMT_INT)
            calc_cell(ws, r, 5, f"=q_W_{code}", FMT_INT)
            calc_cell(ws, r, 6, f"=q_H_{code}", FMT_INT)
            calc_cell(ws, r, 7, f"=INDEX(in_eng_k_len,1,C{r})", "0.00")
            calc_cell(ws, r, 8, f"=INDEX(in_eng_w_allow,1,C{r})", "0")
            calc_cell(ws, r, 9, f"=2*D{r}+G{r}*F{r}", "0")
            calc_cell(ws, r, 10, f"=E{r}+2*F{r}+H{r}", "0")
            calc_cell(ws, r, 11, f"=I{r}*J{r}/1000000", "0.0000")
            calc_cell(ws, r, 12, f"=INDEX(in_eng_net_area,1,C{r})", "0.00")
            calc_cell(ws, r, 13, f"=INDEX(in_eng_grammage,1,C{r})", "0")
            calc_cell(ws, r, 14, f"=K{r}*L{r}*M{r}/1000", "0.0000")
            calc_cell(ws, r, 15, f"=INDEX(in_eng_wrap_mm,1,C{r})", "0")
            calc_cell(ws, r, 16, f"=I{r}+O{r}", "0")
            calc_cell(ws, r, 17, f"=J{r}+O{r}", "0")
            calc_cell(ws, r, 18, f"=INDEX(in_eng_per_bundle,1,C{r})", "0")
            calc_cell(ws, r, 19, f"=INDEX(in_eng_pitch,1,C{r})", "0.00")
            calc_cell(ws, r, 20, f"=R{r}*S{r}", "0.0")
            calc_cell(ws, r, 21, f"=R{r}*N{r}+INDEX(in_eng_wrap_kg,1,C{r})", "0.000")
            calc_cell(ws, r, 22, f"=P{r}*Q{r}*T{r}/1000000000", "0.00000")
            calc_cell(ws, r, 23, f"=V{r}/R{r}", "0.0000000")
            calc_cell(ws, r, 24, f"=U{r}/R{r}", "0.00000")
            r += 1
        rl = r - 1
        for nm, c in [("blankL", "I"), ("blankW", "J"), ("blankkg", "N"), ("a", "P"), ("b", "Q"), ("n", "R"), ("h", "T"), ("bkg", "U"),
                      ("bvol", "V"), ("vpu", "W"), ("kpu", "X")]:
            book.name(f"lg_{nm}_{code}", ws.title, f"{c}{r_first}:{c}{rl}")
        geoA[code] = r_first
    r += 1
    # ---------------- Block B: usable container envelope
    r = section(ws, r, "B. Usable container envelope by scenario", 24)
    r = header(ws, r, ["Ctr", "Scenario", "Scen#", "Inside L", "Inside W", "Inside H", "Door H", "Len ded", "Wid ded", "Top clr", "Usable L", "Usable W",
                       "Usable H", "Road kg", "Cube m³", "Efficiency"])
    for ci, ck in enumerate(CT_KEYS, 1):
        r_first = r
        for si, sn in enumerate(I.SCENARIOS, 1):
            put(ws, r, 1, ck, bold=True)
            put(ws, r, 2, sn)
            put(ws, r, 3, si)
            calc_cell(ws, r, 4, f"=INDEX(ct_L,{ci})", FMT_INT)
            calc_cell(ws, r, 5, f"=INDEX(ct_W,{ci})", FMT_INT)
            calc_cell(ws, r, 6, f"=INDEX(ct_H,{ci})", FMT_INT)
            calc_cell(ws, r, 7, f"=INDEX(ct_doorH,{ci})", FMT_INT)
            calc_cell(ws, r, 8, f"=INDEX(in_use_len_ded,1,C{r})", FMT_INT)
            calc_cell(ws, r, 9, f"=INDEX(in_use_wid_ded,1,C{r})", FMT_INT)
            calc_cell(ws, r, 10, f"=INDEX(in_use_hgt_ded,1,C{r})", FMT_INT)
            calc_cell(ws, r, 11, f"=D{r}-H{r}", FMT_INT)
            calc_cell(ws, r, 12, f"=E{r}-I{r}", FMT_INT)
            calc_cell(ws, r, 13, f"=MIN(F{r}-J{r},G{r})", FMT_INT)
            calc_cell(ws, r, 14, f"=INDEX(ct_road,{ci})", FMT_INT)
            calc_cell(ws, r, 15, f"=INDEX(ct_cube,{ci})", "0.0")
            calc_cell(ws, r, 16, f"=INDEX(in_use_efficiency,1,C{r})", FMT_PCT)
            r += 1
        rl = r - 1
        for nm, c in [("Lu", "K"), ("Wu", "L"), ("Hu", "M"), ("road", "N"), ("cube", "O")]:
            book.name(f"lg_{nm}_{ck}", ws.title, f"{c}{r_first}:{c}{rl}")
    r += 1
    # ---------------- Block C: capacity
    r = section(ws, r, "C. Single-SKU capacity per container (row-mix floor pattern; exact guillotine DP shown for validation)", 34)
    labels = ["SKU", "Ctr", "Scenario", "Scen#", "a mm", "b mm", "h mm", "Bundle kg", "Bundle m³", "Blanks/bdl", "Usable L", "Usable W", "Usable H",
              "Layers", "Pos n2=0", "Pos n2=1", "Pos n2=2", "Floor positions", "Best n2", "Theo bundles", "Theo units", "Efficiency", "Oper. bundles",
              "Oper. units", "Road kg", "Weight-limit bdl", "Weight-limit units", "REALISTIC units", "Binding constraint", "Cargo kg", "Cube util",
              "Payload util", "DP positions (script)", "Validation"]
    r = header(ws, r, labels)
    capC = {}
    for k in I.SKUS:
        code = k["code"]
        for ci, ck in enumerate(CT_KEYS, 1):
            r_first = r
            for si, sn in enumerate(I.SCENARIOS, 1):
                put(ws, r, 1, code, bold=True)
                put(ws, r, 2, ck)
                put(ws, r, 3, sn)
                put(ws, r, 4, si)
                calc_cell(ws, r, 5, f"=INDEX(lg_a_{code},D{r})", "0")
                calc_cell(ws, r, 6, f"=INDEX(lg_b_{code},D{r})", "0")
                calc_cell(ws, r, 7, f"=INDEX(lg_h_{code},D{r})", "0.0")
                calc_cell(ws, r, 8, f"=INDEX(lg_bkg_{code},D{r})", "0.000")
                calc_cell(ws, r, 9, f"=INDEX(lg_bvol_{code},D{r})", "0.00000")
                calc_cell(ws, r, 10, f"=INDEX(lg_n_{code},D{r})", "0")
                calc_cell(ws, r, 11, f"=INDEX(lg_Lu_{ck},D{r})", FMT_INT)
                calc_cell(ws, r, 12, f"=INDEX(lg_Wu_{ck},D{r})", FMT_INT)
                calc_cell(ws, r, 13, f"=INDEX(lg_Hu_{ck},D{r})", FMT_INT)
                calc_cell(ws, r, 14, f"=INT(M{r}/G{r})", "0")
                calc_cell(ws, r, 15, f"=INT(L{r}/F{r})*INT(K{r}/E{r})", "0")
                calc_cell(ws, r, 16, f"=IF(L{r}-E{r}>=0,INT((L{r}-E{r})/F{r})*INT(K{r}/E{r})+INT(K{r}/F{r}),0)", "0")
                calc_cell(ws, r, 17, f"=IF(L{r}-2*E{r}>=0,INT((L{r}-2*E{r})/F{r})*INT(K{r}/E{r})+2*INT(K{r}/F{r}),0)", "0")
                calc_cell(ws, r, 18, f"=MAX(O{r}:Q{r})", "0", bold=True)
                calc_cell(ws, r, 19, f"=MATCH(R{r},O{r}:Q{r},0)-1", "0")
                calc_cell(ws, r, 20, f"=R{r}*N{r}", FMT_INT)
                calc_cell(ws, r, 21, f"=T{r}*J{r}", FMT_INT)
                calc_cell(ws, r, 22, f"=INDEX(in_use_efficiency,1,D{r})", FMT_PCT)
                calc_cell(ws, r, 23, f"=INT(T{r}*V{r})", FMT_INT)
                calc_cell(ws, r, 24, f"=W{r}*J{r}", FMT_INT)
                calc_cell(ws, r, 25, f"=INDEX(lg_road_{ck},D{r})", FMT_INT)
                calc_cell(ws, r, 26, f"=INT(Y{r}/H{r})", FMT_INT)
                calc_cell(ws, r, 27, f"=Z{r}*J{r}", FMT_INT)
                calc_cell(ws, r, 28, f"=MIN(X{r},AA{r})", FMT_INT, bold=True)
                calc_cell(ws, r, 29, f'=IF(X{r}<=AA{r},"Volume / geometry","Weight (US road limit)")')
                calc_cell(ws, r, 30, f"=AB{r}/J{r}*H{r}", FMT_INT)
                calc_cell(ws, r, 31, f"=AB{r}/J{r}*I{r}/INDEX(lg_cube_{ck},D{r})", FMT_PCT)
                calc_cell(ws, r, 32, f"=AD{r}/Y{r}", FMT_PCT)
                static_cell(ws, r, 33, static["dp"][(code, ck, si - 1)], "0")
                calc_cell(ws, r, 34, f'=IF(AG{r}>R{r},"WARN – DP finds "&(AG{r}-R{r})&" more positions",IF(AG{r}=R{r},"OK – pattern = DP optimum","OK – DP below pattern (check)"))')
                r += 1
            rl = r - 1
            book.name(f"lg_real_{code}_{ck}", ws.title, f"AB{r_first}:AB{rl}")
            book.name(f"lg_layers_{code}_{ck}", ws.title, f"N{r_first}:N{rl}")
            book.name(f"lg_pos_{code}_{ck}", ws.title, f"R{r_first}:R{rl}")
            capC[(code, ck)] = r_first
    status_cf(ws, f"AH{r - 27}:AH{r - 1}")
    book.name("lg_validation", ws.title, f"AH{r - 27}:AH{r - 1}")
    r += 1
    r = section(ws, r, "D. Loading schematics (top view, BASE scenario, produced by scripts/loading_model.py – illustrative, not a verified load plan)", 24)
    for img_name in static["images"]:
        p = os.path.join(PROJ, "outputs", img_name)
        if os.path.exists(p):
            img = XLImage(p)
            img.width = int(img.width * 0.62)
            img.height = int(img.height * 0.62)
            ws.add_image(img, f"A{r + 1}")
            r += int(img.height / 20) + 3
    return capC


def build_shipment(book, static):
    ws = book.ws(SH["ship"])
    title(ws, "Shipment Configuration – quoted order vs optimised full-container options",
          "Units per SKU by scenario; feasibility tested with formula zone layouts (each SKU in its own block of container length). Orange = maximum equal-mix quantity solved by scripts/loading_model.py (re-run if geometry inputs change; the FITS columns re-test it).")
    widths = {c: 9 for c in "ABCDEFGHIJKLMNOPQRSTUVWXYZ"}
    widths.update({"A": 6, "B": 44, "C": 7, "D": 6, "E": 6, "F": 9, "G": 7})
    set_widths(ws, widths)
    ncfg = len(I.CONFIGS)
    # layout constants
    rA = 6
    r = section(ws, 4, "A. Configurations (rows = configuration index 1–10)", 30)
    labels = ["#", "Configuration", "Ctr", "Ctr#", "Qty ctr", "Kind", "SKU"]
    for code in SKU_CODES:
        labels += [f"{code} LOW", f"{code} BASE", f"{code} STRESS"]
    labels += ["Total LOW", "Total BASE", "Total STRESS", "FITS LOW", "FITS BASE", "FITS STRESS", "Len util LOW", "Len util BASE", "Len util STRESS"]
    header(ws, 5, labels)
    first = rA
    # Block D location (computed below) – we need row numbers; compute positions deterministically
    container_cfgs = [i for i, c in enumerate(I.CONFIGS) if c["ctype"] != "LCL"]
    rC0 = first + ncfg + 14          # zone calc header row
    nC = len(container_cfgs) * 3 * 3
    rD0 = rC0 + nC + 4               # per cfg×scen totals header row
    dmap = {}                         # (cfg_idx0, scen) -> row in block D
    rr = rD0 + 1
    for ci in container_cfgs:
        for si in range(1, 4):
            dmap[(ci, si)] = rr
            rr += 1
    for i, c in enumerate(I.CONFIGS):
        rr = first + i
        put(ws, rr, 1, i + 1, bold=True)
        put(ws, rr, 2, f"{c['key']} {c['name']}")
        put(ws, rr, 3, c["ctype"])
        inp_cell(ws, rr, 4, CT_IDX[c["ctype"]], "0")
        inp_cell(ws, rr, 5, c["n"], "0")
        put(ws, rr, 6, c["kind"])
        put(ws, rr, 7, c.get("sku", ""))
        for j, code in enumerate(SKU_CODES):
            for si in range(3):
                cc = 8 + j * 3 + si
                if c["kind"] == "fixed":
                    inp_cell(ws, rr, cc, c["units"][code], FMT_INT)
                elif c["kind"] == "fullmix":
                    static_cell(ws, rr, cc, static["fullmix"][(i, si)], FMT_INT)
                else:
                    if code == c["sku"]:
                        calc_cell(ws, rr, cc, f"=INDEX(lg_real_{code}_{c['ctype']},{si + 1})", FMT_INT)
                    else:
                        calc_cell(ws, rr, cc, 0, FMT_INT)
        for si in range(3):
            calc_cell(ws, rr, 17 + si, f"={col(8 + si)}{rr}+{col(11 + si)}{rr}+{col(14 + si)}{rr}", FMT_INT, bold=True)
            if c["ctype"] == "LCL":
                calc_cell(ws, rr, 20 + si, "=TRUE")
                calc_cell(ws, rr, 23 + si, '=""')
            else:
                dr = dmap[(i, si + 1)]
                calc_cell(ws, rr, 20 + si, f"=J{dr}")
                calc_cell(ws, rr, 23 + si, f"=K{dr}", FMT_PCT)
    last = first + ncfg - 1
    ws.conditional_formatting.add(f"T{first}:V{last}", CellIsRule(operator="equal", formula=["FALSE"], fill=FILL_BAD))
    ws.conditional_formatting.add(f"T{first}:V{last}", CellIsRule(operator="equal", formula=["TRUE"], fill=FILL_OK))
    book.name("cfg_key", ws.title, f"A{first}:A{last}")
    book.name("cfg_name", ws.title, f"B{first}:B{last}")
    book.name("cfg_ct", ws.title, f"D{first}:D{last}")
    book.name("cfg_n", ws.title, f"E{first}:E{last}")
    book.name("cfg_u_3464", ws.title, f"H{first}:J{last}")
    book.name("cfg_u_3463", ws.title, f"K{first}:M{last}")
    book.name("cfg_u_3460", ws.title, f"N{first}:P{last}")
    book.name("cfg_total", ws.title, f"Q{first}:S{last}")
    book.name("cfg_fits", ws.title, f"T{first}:V{last}")
    book.name("cfg_lenutil", ws.title, f"W{first}:Y{last}")
    # ---------------- Block B: answers
    r = last + 2
    r = section(ws, r, "B. Physical-feasibility answers (formula)", 12)
    qa = [
        ("Quoted 60,000 mix fits ONE 40' HC (LOW / BASE / STRESS)", "=IF(INDEX(cfg_fits,1,1),\"Yes\",\"No\")&\" / \"&IF(INDEX(cfg_fits,1,2),\"Yes\",\"No\")&\" / \"&IF(INDEX(cfg_fits,1,3),\"Yes\",\"No\")"),
        ("Quoted 60,000 mix fits ONE 40' GP (LOW / BASE / STRESS)", "=IF(INDEX(cfg_fits,2,1),\"Yes\",\"No\")&\" / \"&IF(INDEX(cfg_fits,2,2),\"Yes\",\"No\")&\" / \"&IF(INDEX(cfg_fits,2,3),\"Yes\",\"No\")"),
        ("Quoted 60,000 mix fits TWO 20' GP (LOW / BASE / STRESS)", "=IF(INDEX(cfg_fits,3,1),\"Yes\",\"No\")&\" / \"&IF(INDEX(cfg_fits,3,2),\"Yes\",\"No\")&\" / \"&IF(INDEX(cfg_fits,3,3),\"Yes\",\"No\")"),
        ("Recommended equipment for the quoted order", "=IF(AND(INDEX(cfg_fits,2,1),INDEX(cfg_fits,2,2),INDEX(cfg_fits,2,3)),\"One 40' GP (fits in all scenarios)\",IF(AND(INDEX(cfg_fits,1,1),INDEX(cfg_fits,1,2),INDEX(cfg_fits,1,3)),\"One 40' HC – a 40' GP fails under STRESS bundle thickness; HC is the robust choice\",\"Split shipment required\"))"),
        ("Length utilisation of 40' HC by quoted order (BASE)", "=TEXT(INDEX(cfg_lenutil,1,2),\"0%\")"),
        ("Max equal-mix units per 40' HC (BASE)", "=TEXT(INDEX(cfg_total,4,2),\"#,##0\")&\" boxes (\"&TEXT(INDEX(cfg_u_3464,4,2),\"#,##0\")&\" per SKU)\""),
    ]
    for lab, f in qa:
        put(ws, r, 2, lab, wrap=True)
        calc_cell(ws, r, 3, f)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=14)
        if lab.startswith("Recommended"):
            book.name("ship_reco", ws.title, f"C{r}")
        r += 1
    # LCL vs FCL breakeven
    put(ws, r, 2, "LCL vs one 20' FCL – volume at which FCL becomes cheaper (active scenario, NY/NJ→Edison)", wrap=True)
    f_fcl = ("(INDEX(in_ori_truck_20,1,sel_scen)/fx_active+INDEX(in_ori_terminal,1,sel_scen)+INDEX(in_ori_othc_20,1,sel_scen)+INDEX(in_ori_seal,1,sel_scen)"
             "+INDEX(tbl_ocean_rate,1,sel_scen)+INDEX(tbl_dest_line,1,sel_scen)*(1+INDEX(in_dra_fsc,1,sel_scen))+INDEX(tbl_dest_tolls,1,sel_scen)"
             "+INDEX(in_dra_chassis_days,1,sel_scen)*INDEX(tbl_dest_chassis,1,sel_scen)+INDEX(in_dra_unload_20,1,sel_scen))")
    f_lclfix = ("(INDEX(in_ori_lcl_truck,1,sel_scen)/fx_active+INDEX(in_ori_lcl_cfs_o,1,sel_scen)+INDEX(in_dra_lcl_cfs_d,1,sel_scen)+INDEX(in_dra_lcl_deliv,1,sel_scen))")
    calc_cell(ws, r, 3, f"=MAX(0,({f_fcl}-{f_lclfix})/(INDEX(in_oce_lcl_rate,1,sel_scen)*INDEX(in_oce_lcl_factor,1,sel_scen)))", "0.0")
    put(ws, r, 4, "m³ of bundles")
    book.name("ship_lcl_be_m3", ws.title, f"C{r}")
    r += 1
    put(ws, r, 2, "… equivalent equal-mix boxes (BASE m³/box)", wrap=True)
    calc_cell(ws, r, 3, "=ship_lcl_be_m3/((INDEX(lg_vpu_3464,sel_scen)+INDEX(lg_vpu_3463,sel_scen)+INDEX(lg_vpu_3460,sel_scen))/3)", FMT_INT)
    put(ws, r, 4, "boxes (LCL rates are ASSUMED – RFQ 2)")
    r += 1
    # ---------------- Block C: zone calc
    r = rC0 - 1
    section(ws, r, "C. Zone-layout calculation per container (each SKU occupies a block of length; candidate-length method, validated against exact search in scripts/loading_model.py)", 32)
    labels = ["Cfg#", "Scen#", "SKU", "Ctr#", "Qty ctr", "Units", "Units/ctr", "Blanks/bdl", "Bundles/ctr", "Layers", "Eff.", "Stacks", "a", "b", "Usable W", "Usable L"]
    for n2 in range(3):
        labels += [f"n1 (n2={n2})", f"i0 (n2={n2})", f"j0 (n2={n2})", f"Lmin (n2={n2})"]
    labels += ["Zone length mm", "Zone kg", "Zone m³"]
    header(ws, rC0, labels)
    rr = rC0 + 1
    zfirst = rr
    for ci in container_cfgs:
        c = I.CONFIGS[ci]
        ck = c["ctype"]
        for si in range(1, 4):
            for code in SKU_CODES:
                put(ws, rr, 1, ci + 1)
                put(ws, rr, 2, si)
                put(ws, rr, 3, code)
                calc_cell(ws, rr, 4, f"=INDEX(cfg_ct,A{rr})", "0")
                calc_cell(ws, rr, 5, f"=INDEX(cfg_n,A{rr})", "0")
                calc_cell(ws, rr, 6, f"=INDEX(cfg_u_{code},A{rr},B{rr})", FMT_INT)
                calc_cell(ws, rr, 7, f"=F{rr}/E{rr}", FMT_INT)
                calc_cell(ws, rr, 8, f"=INDEX(lg_n_{code},B{rr})", "0")
                calc_cell(ws, rr, 9, f"=ROUNDUP(G{rr}/H{rr},0)", FMT_INT)
                calc_cell(ws, rr, 10, f"=INDEX(lg_layers_{code}_{ck},B{rr})", "0")
                calc_cell(ws, rr, 11, f"=INDEX(in_use_efficiency,1,B{rr})", FMT_PCT)
                calc_cell(ws, rr, 12, f"=IF(I{rr}=0,0,ROUNDUP(I{rr}/(K{rr}*J{rr})-0.000000001,0))", "0")
                calc_cell(ws, rr, 13, f"=INDEX(lg_a_{code},B{rr})", "0")
                calc_cell(ws, rr, 14, f"=INDEX(lg_b_{code},B{rr})", "0")
                calc_cell(ws, rr, 15, f"=INDEX(lg_Wu_{ck},B{rr})", "0")
                calc_cell(ws, rr, 16, f"=INDEX(lg_Lu_{ck},B{rr})", "0")
                lcols = []
                for n2 in range(3):
                    c0 = 17 + n2 * 4
                    n1c, i0c, j0c, lmc = col(c0), col(c0 + 1), col(c0 + 2), col(c0 + 3)
                    calc_cell(ws, rr, c0, f"=IF(O{rr}-{n2}*M{rr}>=0,INT((O{rr}-{n2}*M{rr})/N{rr}),-1)", "0")
                    calc_cell(ws, rr, c0 + 1, f"=IF(AND({n1c}{rr}>=0,{n1c}{rr}+{n2}>0),ROUNDUP(L{rr}/({n1c}{rr}+{n2}*M{rr}/N{rr}),0),0)", "0")
                    calc_cell(ws, rr, c0 + 2, f"=IF(AND({n1c}{rr}>=0,{n1c}{rr}+{n2}>0),ROUNDUP(L{rr}/({n1c}{rr}*N{rr}/M{rr}+{n2}),0),0)", "0")
                    cands = []
                    for dlt in (-1, 0, 1, 2):
                        La = f"MAX({i0c}{rr}+({dlt}),0)*M{rr}"
                        cands.append(f"IF(AND({La}>0,{n1c}{rr}*INT({La}/M{rr})+{n2}*INT({La}/N{rr})>=L{rr}),{La},1000000000)")
                    for dlt in (-1, 0, 1, 2):
                        Lb = f"MAX({j0c}{rr}+({dlt}),0)*N{rr}"
                        cands.append(f"IF(AND({Lb}>0,{n1c}{rr}*INT({Lb}/M{rr})+{n2}*INT({Lb}/N{rr})>=L{rr}),{Lb},1000000000)")
                    calc_cell(ws, rr, c0 + 3, f"=IF(OR({n1c}{rr}<0,{n1c}{rr}+{n2}=0),1000000000,MIN({','.join(cands)}))", "0")
                    lcols.append(f"{lmc}{rr}")
                calc_cell(ws, rr, 29, f"=IF(L{rr}=0,0,MIN({','.join(lcols)}))", FMT_INT, bold=True)
                calc_cell(ws, rr, 30, f"=I{rr}*INDEX(lg_bkg_{code},B{rr})", FMT_INT)
                calc_cell(ws, rr, 31, f"=I{rr}*INDEX(lg_bvol_{code},B{rr})", "0.00")
                rr += 1
    zlast = rr - 1
    # ---------------- Block D
    section(ws, rD0 - 1, "D. Feasibility per configuration and scenario (one container; multi-container configs split units evenly)", 14)
    header(ws, rD0, ["Cfg#", "Scen#", "Ctr#", "Σ zone length", "Usable L", "Fits length", "kg / ctr", "Road kg", "Fits weight", "FITS", "Length util", "Cube util", "Payload util"])
    for (ci, si), dr in dmap.items():
        put(ws, dr, 1, ci + 1)
        put(ws, dr, 2, si)
        calc_cell(ws, dr, 3, f"=INDEX(cfg_ct,A{dr})", "0")
        calc_cell(ws, dr, 4, f"=SUMIFS($AC${zfirst}:$AC${zlast},$A${zfirst}:$A${zlast},A{dr},$B${zfirst}:$B${zlast},B{dr})", FMT_INT)
        calc_cell(ws, dr, 5, f"=INDEX(lg_Lu_{I.CONFIGS[ci]['ctype']},B{dr})", FMT_INT)
        calc_cell(ws, dr, 6, f"=D{dr}<=E{dr}+0.000001")
        calc_cell(ws, dr, 7, f"=SUMIFS($AD${zfirst}:$AD${zlast},$A${zfirst}:$A${zlast},A{dr},$B${zfirst}:$B${zlast},B{dr})", FMT_INT)
        calc_cell(ws, dr, 8, f"=INDEX(ct_road,C{dr})", FMT_INT)
        calc_cell(ws, dr, 9, f"=G{dr}<=H{dr}")
        calc_cell(ws, dr, 10, f"=AND(F{dr},I{dr})", bold=True)
        calc_cell(ws, dr, 11, f"=D{dr}/E{dr}", FMT_PCT)
        calc_cell(ws, dr, 12, f"=SUMIFS($AE${zfirst}:$AE${zlast},$A${zfirst}:$A${zlast},A{dr},$B${zfirst}:$B${zlast},B{dr})/INDEX(ct_cube,C{dr})", FMT_PCT)
        calc_cell(ws, dr, 13, f"=G{dr}/H{dr}", FMT_PCT)
    dl = max(dmap.values())
    ws.conditional_formatting.add(f"J{rD0 + 1}:J{dl}", CellIsRule(operator="equal", formula=["FALSE"], fill=FILL_BAD))
    ws.conditional_formatting.add(f"J{rD0 + 1}:J{dl}", CellIsRule(operator="equal", formula=["TRUE"], fill=FILL_OK))
    book.name("ship_zone_len", ws.title, f"AC{zfirst}:AC{zlast}")
    # static check: script fullmix still fits
    book.name("ship_fullmix_fits", ws.title, f"T{first + 3}:V{first + 8}")


def build_origin(book):
    ws = book.ws(SH["orig"])
    title(ws, "Origin Costs – Egypt (from the supplier's premises to vessel loading)",
          "Inputs by scenario. EGP items convert at the scenario's effective FX. The selected-case ledger is on 'Landed Cost by SKU'.")
    set_widths(ws, {"A": 9, "B": 52, "C": 11, "D": 11, "E": 11, "F": 11, "G": 14, "H": 18, "I": 16, "J": 70})
    r = section(ws, 4, "1. Origin cost inputs", 10)
    rows, r = input_table(book, ws, r, "ori", list(I.ORIGIN.items()), "ORI")
    r += 1
    r = section(ws, r, "2. Selected case – origin cost summary (formula)", 10)
    r = header(ws, r, ["", "Item", "", "USD", "Per box", "", "", "", "", "Note"])
    lines = [("Per-container origin costs (trucking, terminal/VGM, THC, seal, inspection, detention, loading)", "=INDEX(eng_orig_ctr,sel_row)"),
             ("Per-shipment origin costs (clearance, CoO, B/L, forwarder, QC, bank wires)", "=INDEX(eng_orig_shp,sel_row)"),
             ("TOTAL ORIGIN COSTS", "=INDEX(eng_ORIGIN,sel_row)")]
    for lab, f in lines:
        put(ws, r, 2, lab, wrap=True, bold=lab.startswith("TOTAL"))
        calc_cell(ws, r, 4, f, FMT_USD, bold=lab.startswith("TOTAL"))
        calc_cell(ws, r, 5, f"=D{r}/INDEX(eng_SALE,sel_row)", FMT_USD4)
        r += 1
    r += 1
    r = section(ws, r, "3. EXW versus FCA – why the working assumption is fragile", 10)
    txt = [
        "EXW (Incoterms 2020): seller only makes goods available at its premises; no obligation to load or to clear for export. Buyer handles export formalities (S072).",
        "Egypt: exports may only be made by parties registered in the Exporters Register (Law 118/1975, Art. 4; S062). A foreign buyer without an Egyptian registered exporter cannot itself clear an EXW purchase for export.",
        "ICC strongly encourages FCA instead of EXW when goods cross a border. Under FCA (named place = NFC factory) the seller loads the container and clears export; the buyer's forwarder takes over at the factory gate.",
        "Recommendation: request FCA [NFC factory, 6th of October] Incoterms 2020 with NFC as exporter of record (VAT zero-rated export invoice), and an EXW price for comparison. Do not silently replace EXW – the model keeps EXW pricing and adds origin costs explicitly.",
        "Customs value (19 CFR 152.103(a)(5)): if the sale is genuinely ex-factory, Egyptian inland freight is not added to customs value; under FCA-factory the same applies if inland freight is invoiced separately by the forwarder.",
    ]
    for t in txt:
        put(ws, r, 2, t, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
        ws.row_dimensions[r].height = 32
        r += 1


def build_ocean(book):
    ws = book.ws(SH["ocean"])
    title(ws, "Ocean Freight & Insurance – Egypt (Alexandria / Damietta) → US East Coast",
          "No written quote was obtained. Rates are ESTIMATED from carrier tariff notices and market indices (see evidence table). Base port rate + port differential.")
    set_widths(ws, {"A": 9, "B": 52, "C": 11, "D": 11, "E": 11, "F": 11, "G": 14, "H": 18, "I": 22, "J": 70})
    r = section(ws, 4, "1. Ocean & insurance inputs (first three rows form the rate table by container: 20', 40', 40' HC)", 10)
    rows, r = input_table(book, ws, r, "oce", list(I.OCEAN.items()), "OCE")
    book.name("tbl_ocean_rate", ws.title, f"D{rows['rate_20']}:F{rows['rate_hc']}")
    r += 1
    r = section(ws, r, "2. Port differential over base port (USD per container)", 10)
    r = header(ws, r, ["#", "Gateway", "", "LOW", "BASE", "STRESS", "ACTIVE", "Evidence", "Source", "Note"])
    pfirst = r
    for i, (pk, pn) in enumerate(PORTS, 1):
        put(ws, r, 1, i)
        put(ws, r, 2, pn)
        for j in range(3):
            inp_cell(ws, r, 4 + j, I.PORT_DIFF[pk][j], FMT_USD0)
        calc_cell(ws, r, 7, f"=INDEX(D{r}:F{r},1,sel_scen)", FMT_USD0)
        put(ws, r, 8, "ESTIMATED" if pk == "SAV" else ("VERIFIED_MARKET" if pk == "NYNJ" else "ASSUMED"))
        put(ws, r, 9, "S053" if pk == "SAV" else ("S047" if pk == "NYNJ" else ""))
        put(ws, r, 10, {"NYNJ": "Base port (CMA CGM FAK quoted to New York)", "BAL": "Fewer direct Med→Baltimore calls; often via transshipment",
                        "ORF": "Usually base-port pricing", "SAV": "Savannah peak-season premium $75–200/TEU (S053)"}[pk], wrap=True)
        r += 1
    book.name("tbl_port_diff", ws.title, f"D{pfirst}:F{r - 1}")
    book.name("port_name", ws.title, f"B{pfirst}:B{r - 1}")
    evidence_cf(ws, f"H{pfirst}:H{r - 1}")
    r += 1
    r = section(ws, r, "3. Market evidence used to build the estimates", 10)
    r = header(ws, r, ["Source", "Benchmark", "Lane", "20'", "40'", "Date", "Type", "", "", "Use in model"])
    ev = [("S047", "CMA CGM FAK (incl. O/D THC; excl. bunker, AMS, ETS, PSS)", "Genoa→New York", 2750, 3600, "from 2026-04-01", "Carrier tariff notice", "BASE anchor + ~$250–300 surcharges"),
          ("S048", "Hapag-Lloyd GRI Egypt→North America", "Egypt→NA", 300, 600, "from 2026-07-01", "Carrier notice (increment)", "STRESS add-on"),
          ("S045", "FBX22 Europe→US East Coast", "N. Europe→USEC", None, 2723, "2026-10-02", "Index (spot)", "LOW anchor"),
          ("S046", "Xeneta N. Europe→USEC", "N. Europe→USEC", None, 2741, "2026-08-12", "Index (spot)", "LOW cross-check"),
          ("S044", "Drewry WCI Rotterdam→New York", "Rotterdam→NY", None, 3116, "2026-10-08", "Index (spot, all-in)", "Cross-check"),
          ("S049", "Globy indicative Alexandria→New York (expired validity)", "Alexandria→NY", 1336, None, "2025–Mar 2026", "Marketplace", "Routing & transit (25–33 days)"),
          ("S095", "CMA CGM US/Canada rate-restoration charges from W. Med", "W. Med→US", None, None, "from 2026-10-01", "Carrier notice", "Supports STRESS")]
    for sid, b, lane, v20, v40, d, t, use in ev:
        put(ws, r, 1, sid)
        put(ws, r, 2, b, wrap=True)
        put(ws, r, 3, lane)
        static_cell(ws, r, 4, v20, FMT_USD0) if v20 else put(ws, r, 4, "–")
        static_cell(ws, r, 5, v40, FMT_USD0) if v40 else put(ws, r, 5, "–")
        put(ws, r, 6, d)
        put(ws, r, 7, t)
        put(ws, r, 10, use, wrap=True)
        r += 1
    r += 1
    r = section(ws, r, "4. Selected case – ocean & insurance (formula)", 10)
    for lab, f, fm in [("Ocean freight (container, or LCL all-in W/M)", "=INDEX(eng_ocean_ctr,sel_row)", FMT_USD),
                       ("AMS / manifest", "=INDEX(in_oce_ams,1,sel_scen)", FMT_USD),
                       ("Cargo insurance (MAX(min, rate × (supplier + origin + ocean) × 1.1))", "=INDEX(eng_INS,sel_row)", FMT_USD),
                       ("Ocean freight per saleable box", "=INDEX(eng_OCEAN,sel_row)/INDEX(eng_SALE,sel_row)", FMT_USD4),
                       ("Transit days (incl. LCL extra if applicable)", "=INDEX(eng_t2,sel_row)-INDEX(eng_t1,sel_row)", "0")]:
        put(ws, r, 2, lab, wrap=True)
        calc_cell(ws, r, 4, f, fm)
        r += 1
    r += 1
    r = section(ws, r, "5. Insurance notes", 10)
    for t in ["Insurance is not included in EXW/FCA or in ocean freight quotes unless stated (double-count flag OCE input).",
              "Insured value basis: CIF + 10% (commercial practice, S089). It is unrelated to US customs value (FOB-type; international freight and insurance are excluded from US transaction value).",
              "Recommended cover: ICC(A) all risks incl. war/strikes; check exclusions for wetting/condensation of paperboard and inadequate packing; deductible and minimum premium to be quoted (G15)."]:
        put(ws, r, 2, t, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
        ws.row_dimensions[r].height = 28
        r += 1


def build_us_customs(book):
    ws = book.ws(SH["usc"])
    title(ws, "U.S. Port & Customs – entry, fees, bond, broker, port charges",
          "Duty and government ad-valorem fees are on 'Tariffs & Customs Value'. One formal entry per shipment assumed (value > $2,500).")
    set_widths(ws, {"A": 9, "B": 52, "C": 14, "D": 11, "E": 11, "F": 11, "G": 14, "H": 18, "I": 14, "J": 70})
    r = section(ws, 4, "1. Entry, bond, port and disruption inputs", 10)
    keys = ["broker", "isf", "stb_rate", "stb_min", "isf_bond", "cfc", "cfc_in_freight", "do_fee", "exam", "exam_dray", "dem_days", "dem_rate"]
    rows, r = input_table(book, ws, r, "cus", [(k, I.CUSTOMS[k]) for k in keys], "USC")
    r += 1
    r = section(ws, r, "2. Selected case (formula)", 10)
    for lab, f in [("Single-transaction bond premium", "=INDEX(eng_STB,sel_row)"), ("Per-shipment US fees (broker + ISF + DO + bond + ISF bond)", "=INDEX(eng_US_SHP,sel_row)"),
                   ("PANYNJ CFC (NY/NJ only)", "=INDEX(eng_CFC,sel_row)"), ("CBP exam + exam drayage (STRESS)", "=INDEX(eng_EXAM,sel_row)"),
                   ("Demurrage / storage (STRESS)", "=INDEX(eng_DEM,sel_row)"), ("MPF", "=INDEX(eng_MPF,sel_row)"), ("HMF", "=INDEX(eng_HMF,sel_row)"),
                   ("Duty (MFN + Section 301)", "=INDEX(eng_DUTY,sel_row)")]:
        put(ws, r, 2, lab)
        calc_cell(ws, r, 4, f, FMT_USD)
        r += 1
    r += 1
    r = section(ws, r, "3. One-time vs per-shipment costs", 10)
    r = header(ws, r, ["", "Item", "Type", "Indicative", "", "", "", "Evidence", "Source", "Note"])
    for item, typ, amt, ev, src, note in [
        ("Importer identity (CBP Form 5106) – IRS EIN or CBP-assigned number", "One-time", "$0 (form)", "VERIFIED_OFFICIAL", "S040", "Foreign entities can obtain a CBP-assigned number"),
        ("Broker onboarding, power of attorney, KYC", "One-time", "$0–250", "ASSUMED", "", ""),
        ("US entity formation (optional – LLC + registered agent)", "One-time / annual", "$500–2,500", "ASSUMED", "", "Not required to be IOR"),
        ("Continuous bond ($50,000 minimum)", "Annual", "$250–600 / yr", "ESTIMATED", "S038", "Cheaper than STB above ~5–10 entries a year"),
        ("Single-transaction bond + ISF bond", "Per shipment", "≥ $75 + $60", "ESTIMATED", "S038", "Used in model"),
        ("Customs entry (broker)", "Per shipment", "$150–300", "ESTIMATED", "S039", ""),
        ("ISF 10+2 filing", "Per shipment", "$35–75", "ESTIMATED", "S039", "≥24 h before loading; $5,000 liquidated damages"),
        ("MPF 0.3464% (min $34.58, max $670.86, FY2027)", "Per entry", "formula", "VERIFIED_OFFICIAL", "S024", ""),
        ("HMF 0.125%", "Per entry (ocean)", "formula", "VERIFIED_OFFICIAL", "S025", ""),
        ("FDA prior notice / food-facility registration", "n/a", "$0", "VERIFIED_OFFICIAL", "S036", "Food-contact substances excluded from 'food' for these rules (read via a secondary summary)"),
    ]:
        put(ws, r, 2, item, wrap=True)
        put(ws, r, 3, typ)
        put(ws, r, 4, amt)
        put(ws, r, 8, ev)
        put(ws, r, 9, src)
        put(ws, r, 10, note, wrap=True)
        r += 1
    r += 1
    r = section(ws, r, "4. First-shipment compliance checklist (importer of record)", 10)
    for t in [
        "Decide IOR: BigBoxPizza (foreign IOR with CBP-assigned number), a US entity, a third-party IOR service, or a US distributor that imports (see Read Me §IOR).",
        "Appoint a licensed customs broker; sign customs power of attorney; file CBP Form 5106.",
        "Bond: single-transaction (pilot) or continuous ($50k) bond; ISF bond if no continuous bond.",
        "ISF 10+2 filed ≥24 h before vessel loading in Egypt (seller, buyer, consignee, manufacturer, ship-to, origin, HTS, stuffing location, consolidator, IOR).",
        "Commercial invoice (seller, buyer, description 'Corrugated paperboard pizza boxes, flat-packed (die-cut, unassembled), sizes…', HTS 4819.10.00, unit & total value, currency, Incoterm, country of origin Egypt), packing list (bundles, weights, dimensions).",
        "Bill of lading (or sea waybill); arrival notice; delivery order.",
        "Classification note: 4819.10.0020 (sanitary food containers) vs .0040 – broker to confirm; Chapter 99: 9903.05.36 (Egypt, 12.5%) unless an exemption applies.",
        "Valuation support: proof of ex-factory/FCA price, payment records, separate freight invoices (foreign inland freight not added for ex-factory sales).",
        "Origin marking: 'Made in Egypt' on each bundle wrap / outer pack reaching the user (19 CFR 134.24) – broker to confirm.",
        "Food-contact file: supplier declaration (21 CFR 176.170/176.260), PFAS-free statement and total-fluorine test, TPCH heavy-metal certificate – required by buyers and by NY/MD law for PFAS.",
        "Wood packaging: if pallets are used they must be ISPM-15 heat-treated and marked (S090).",
        "Arrange drayage and delivery appointment before arrival; confirm free time; pay duties/fees via broker ACH before release.",
    ]:
        put(ws, r, 2, "☐ " + t, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
        ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    r = section(ws, r, "5. Importer-of-record structures (Workstream L) – qualitative comparison", 10)
    r = header(ws, r, ["", "Structure", "Legal feasibility", "IOR liability", "Cash-flow", "Customer control", "Margin kept", "Complexity", "Set-up / recurring", "Key issues"])
    for row in [
        ("BigBoxPizza as foreign IOR", "Feasible (CBP-assigned number + bond + broker)", "BigBoxPizza", "Pays duty/fees before release", "Full", "Full", "Medium", "Low / per-entry", "Needs US broker, bond, possibly US tax advice; liability for classification/value"),
        ("New US entity (LLC) as IOR", "Feasible", "US entity", "Same; US banking", "Full", "Full", "High", "$500–2,500 + annual", "US tax filings; better for repeat trade"),
        ("Third-party IOR service", "Feasible", "Service provider (contractual)", "Provider fees", "Full", "Lower (fees)", "Low", "Per-shipment fee", "Check that the provider genuinely assumes IOR liability; broker ≠ IOR"),
        ("US distributor imports (sell FCA/FOB Egypt)", "Feasible", "Distributor", "Distributor pays freight/duty", "Shared", "Lowest", "Low", "Minimal", "Buyer bears tariff/logistics risk; BigBoxPizza becomes an export trader"),
    ]:
        for j, v in enumerate(row):
            put(ws, r, 2 + j, v, wrap=True)
        ws.row_dimensions[r].height = 42
        r += 1
    put(ws, r + 1, 2, "A customs broker files entries as agent; it does not assume the importer's legal obligations (reasonable care, correct classification and value).", wrap=True)
    ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=10)


def build_destination(book):
    ws = book.ws(SH["dest"])
    title(ws, "Destination Delivery – gateway → representative customer location",
          "Distance-band ESTIMATES (no written trucking quotes). Customer locations are representative ZIP areas, not real customers.")
    set_widths(ws, {"A": 6, "B": 22, "C": 52, "D": 8, "E": 9, "F": 9, "G": 9, "H": 9, "I": 9, "J": 9, "K": 9, "L": 9, "M": 9, "N": 7, "O": 14, "P": 14, "Q": 14})
    r = section(ws, 4, "1. Destinations (rows = destination index 1–5)", 17)
    r = header(ws, r, ["#", "Gateway", "Representative delivery location", "Miles", "Line LOW", "Line BASE", "Line STRESS", "Tolls LOW", "Tolls BASE",
                       "Tolls STRESS", "Chassis $/d LOW", "Chassis BASE", "Chassis STRESS", "Port#", "Per-ctr cost LOW", "BASE", "STRESS"])
    dfirst = r
    for i, d in enumerate(I.DESTINATIONS, 1):
        put(ws, r, 1, i, bold=True)
        put(ws, r, 2, d["port_name"])
        put(ws, r, 3, d["city"], wrap=True)
        inp_cell(ws, r, 4, d["miles"], "0")
        for j in range(3):
            inp_cell(ws, r, 5 + j, d["line"][j], FMT_USD0)
            inp_cell(ws, r, 8 + j, d["tolls"][j], FMT_USD0)
            inp_cell(ws, r, 11 + j, d["chassis"][j], FMT_USD0)
        inp_cell(ws, r, 14, PORT_IDX[d["port"]], "0")
        for j in range(3):
            sc = j + 1
            f = (f"={col(5 + j)}{r}*(1+INDEX(in_dra_fsc,1,{sc}))+{col(8 + j)}{r}+INDEX(in_dra_chassis_days,1,{sc})*{col(11 + j)}{r}*(1-INDEX(in_dra_chassis_in_dray,1,{sc}))"
                 f"+INDEX(in_dra_wait_hours,1,{sc})*INDEX(in_dra_wait_rate,1,{sc})+INDEX(in_dra_split,1,{sc})+INDEX(in_dra_prepull,1,{sc})+INDEX(in_dra_dryrun,1,{sc})"
                 f"+INDEX(in_dra_unload_40,1,{sc})*(1-INDEX(in_dra_unload_in_dray,1,{sc}))")
            calc_cell(ws, r, 15 + j, f, FMT_USD0)
        r += 1
    dl = r - 1
    book.name("tbl_dest_line", ws.title, f"E{dfirst}:G{dl}")
    book.name("tbl_dest_tolls", ws.title, f"H{dfirst}:J{dl}")
    book.name("tbl_dest_chassis", ws.title, f"K{dfirst}:M{dl}")
    book.name("dest_port", ws.title, f"N{dfirst}:N{dl}")
    book.name("dest_name", ws.title, f"C{dfirst}:C{dl}")
    book.name("dest_gw", ws.title, f"B{dfirst}:B{dl}")
    put(ws, r, 3, "Per-ctr cost columns = 40' delivery incl. FSC, tolls, chassis, waiting, accessorials and unloading (for comparison). Sources: " + I.DEST_SOURCES, wrap=True)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=17)
    r += 2
    r = section(ws, r, "2. Drayage, accessorial and unloading inputs", 10)
    rows, r = input_table(book, ws, r, "dra", list(I.DRAY.items()), "DRA")
    r += 1
    r = section(ws, r, "3. Selected case (formula)", 10)
    for lab, f in [("Delivery (drayage incl. accessorials, or LCL destination fixed charges + delivery)", "=INDEX(eng_DELIV,sel_row)"), ("Unloading", "=INDEX(eng_UNLOAD,sel_row)"),
                   ("Delivery + unloading per saleable box", "=(INDEX(eng_DELIV,sel_row)+INDEX(eng_UNLOAD,sel_row))/INDEX(eng_SALE,sel_row)")]:
        put(ws, r, 3, lab)
        calc_cell(ws, r, 5, f, FMT_USD4 if "per" in lab else FMT_USD)
        r += 1
    r += 1
    r = section(ws, r, "4. Delivery model notes", 10)
    for t in ["Principal model: full container drayed from the port terminal directly to the wholesale customer's dock, live-unloaded (floor-loaded bundles), empty returned to the terminal – no long-term storage.",
              "Floor-loaded flat bundles (≈1,200–1,600 bundles, 10–15 t per 40') must be hand-unloaded; most drayage quotes exclude unloading (flag input). Waiting time beyond 2 free hours is modelled.",
              "Port arrival cost is NOT the landed cost: customer-delivered cost adds drayage, chassis, tolls, waiting, unloading (and disruption items in STRESS).",
              "Disruption scenario (STRESS): CBP exam + exam drayage, 3 days demurrage/storage, pre-pull & yard storage, dry run, 5 chassis days, 4 waiting hours.",
              "Savannah→Atlanta could alternatively use CSX rail to the Appalachian Regional Port (S094 context) – not modelled."]:
        put(ws, r, 3, t, wrap=True)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=17)
        ws.row_dimensions[r].height = 30
        r += 1


def build_tariffs(book):
    ws = book.ws(SH["tariff"])
    title(ws, "Tariffs & Customs Value – classification, duty layers, customs valuation",
          "Statutory rate ≠ applicable treatment ≠ calculated amount ≠ quoted amount ≠ final amount paid. Broker confirmation required (RFQ 3).")
    set_widths(ws, {"A": 9, "B": 46, "C": 14, "D": 12, "E": 12, "F": 12, "G": 14, "H": 20, "I": 18, "J": 60})
    r = section(ws, 4, "1. Duty and government-fee inputs", 10)
    keys = ["mfn", "s301", "mpf_rate", "mpf_min", "mpf_max", "hmf"]
    rows, r = input_table(book, ws, r, "cus", [(k, I.CUSTOMS[k]) for k in keys], "TAR")
    r += 1
    r = section(ws, r, "2. Classification candidates (HTSUS, 2026 Rev. 13 via mirror)", 10)
    r = header(ws, r, ["Rank", "Heading / description", "HTS", "MFN (General)", "Chapter 99 (Egypt)", "", "Confidence", "Evidence", "Sources", "Reasoning"])
    for row in [
        ("1", "Cartons, boxes and cases, of corrugated paper or paperboard – stat. suffix .0020 'Sanitary food and beverage containers' (likely) or .0040 'Other'", "4819.10.00", "Free", "9903.05.36 +12.5%", "High (heading); Medium (suffix)", "VERIFIED_OFFICIAL", "S010;S011;S013;S014;S016",
         "Corrugated E-flute pizza boxes are corrugated boxes; flat die-cut blanks imported unassembled are classified as the complete box under GRI 2(a) (CBP practice for knocked-down boxes)."),
        ("2", "Folding cartons of NON-corrugated paperboard", "4819.20.00", "Free", "+12.5%", "Low", "ESTIMATED", "S013",
         "Only if 'SINGLE-E' were not corrugated (unlikely) – resolve with specification."),
        ("3", "Other articles of paper (residual)", "4823.90.86", "Free", "+12.5%", "Low", "VERIFIED_MARKET", "S101", "Residual heading; rejected because 4819 specifically describes boxes."),
        ("4", "Corrugated paperboard in sheets", "4808.10.00", "Free", "+12.5%", "Low", "VERIFIED_MARKET", "S101", "Rejected: die-cut, scored blanks for a box have the essential character of the box."),
    ]:
        for j, v in enumerate(row):
            put(ws, r, 1 + j if j < 5 else 2 + j, v, wrap=True)
        ws.row_dimensions[r].height = 54
        r += 1
    put(ws, r, 2, "Conclusion: all plausible headings are MFN Free; the economically relevant layer is the Egypt Section 301 duty, which applies regardless of heading unless an Annex II exemption is confirmed.", wrap=True, bold=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
    ws.row_dimensions[r].height = 30
    r += 2
    r = section(ws, r, "3. Duty layers and trade measures checked (as of research date)", 10)
    r = header(ws, r, ["", "Measure", "Legal authority", "Effective", "Applies?", "Rate in model", "", "Evidence", "Sources", "Notes"])
    for m in [
        ("Ordinary (MFN) duty, HTS 4819.10.00", "HTSUS column 1", "current", "Yes", "Free", "VERIFIED_OFFICIAL", "S010;S011", ""),
        ("Section 301 forced-labor action – Egypt (9903.05.36)", "Trade Act 1974 §301; USTR final action FR 2026-15274/15181; CSMS 69326983", "2026-07-24", "Probable (no exemption found)", "12.5% (LOW 0%)", "VERIFIED_OFFICIAL", "S014;S015;S016;S017;S018", "Exemptions: Annex II Part A (2,120 codes; energy, critical minerals, precious metals, certain agri inputs, aircraft/pharma use), §232 goods, Ch.98, USMCA. CIT challenge argued 30-Sep-2026."),
        ("IEEPA 'reciprocal' 10% on Egypt", "IEEPA – struck down", "ended 2026-02-23", "No", "0%", "VERIFIED_OFFICIAL", "S021", "Learning Resources v. Trump (20-Feb-2026)"),
        ("Section 122 global surcharge 10%", "Trade Act 1974 §122", "2026-02-24 to 2026-07-24", "No (expired)", "0%", "VERIFIED_MARKET", "S022", "Lapsed at 150-day limit"),
        ("Section 232 (timber/lumber & derivatives)", "Proclamation 10976", "2025-10", "No (ch. 48 not covered per summaries)", "0%", "VERIFIED_MARKET", "S023", ""),
        ("AD/CVD – corrugated die-cut (pizza) boxes", "Petitions 9-Sep-2026 (China, Malaysia, Türkiye)", "pending", "No (Egypt not named)", "0%", "VERIFIED_OFFICIAL", "S097;S098;S100", "Risk if Egyptian imports surge"),
        ("QIZ duty-free (US–Israel FTA authority)", "Proclamation 6955 / QIZ protocol", "n/a", "Unproven", "not used", "UNRESOLVED", "S071", "Requires QIZ-registered producer and 10.5% Israeli content; interaction with §301 unconfirmed"),
    ]:
        for j, v in enumerate(m):
            put(ws, r, 2 + j if j < 5 else 3 + j, v, wrap=True)
        ws.row_dimensions[r].height = 48
        r += 1
    evidence_cf(ws, f"H{r - 7}:H{r - 1}")
    r += 1
    r = section(ws, r, "4. Customs value and duty – selected case (formula, per SKU)", 10)
    r = header(ws, r, ["", "Element", "3464", "3463", "3460", "TOTAL", "", "Treatment", "Authority", "Notes"])
    lines = [
        ("Price actually paid – EXW/FCA product value (USD)", ["=INDEX(eng_sup_3464,sel_row)", "=INDEX(eng_sup_3463,sel_row)", "=INDEX(eng_sup_3460,sel_row)"], "Included", "19 USC 1401a; 19 CFR 152.103", ""),
        ("+ Egyptian VAT if paid and NOT refunded (mode 3)", [f"=IF(INDEX(eng_vatmode,sel_row)=3,INDEX(eng_sup_{c},sel_row)*INDEX(in_sup_vat_rate,1,sel_scen),0)" for c in SKU_CODES], "Conditional", "Price actually paid", "Refunded/zero-rated export VAT excluded"),
        ("+ Tooling / dies paid to seller (allocated by units)", [f"=INDEX(eng_tooling,sel_row)*INDEX(eng_units_{c},sel_row)/INDEX(eng_U,sel_row)" for c in SKU_CODES], "Included (assist/price)", "19 USC 1401a(b)(1)", ""),
        ("+ Packing costs incurred by buyer", ["=0", "=0", "=0"], "Included if any", "19 USC 1401a(b)(1)(A)", "Bundling assumed in price"),
        ("− Egyptian inland freight (ex-factory sale)", ["=0", "=0", "=0"], "Not added", "19 CFR 152.103(a)(5)(i)", "Must prove ex-factory sale"),
        ("− International freight & insurance", ["=0", "=0", "=0"], "Excluded", "19 USC 1401a(b)(4)(A)", "Not part of US value"),
        ("CUSTOMS (DUTIABLE) VALUE", [f"=INDEX(eng_cv_{c},sel_row)" for c in SKU_CODES], "", "", "Check equals sum above"),
        ("MFN duty (Free)", [f"=INDEX(eng_cv_{c},sel_row)*INDEX(in_cus_mfn,1,sel_scen)" for c in SKU_CODES], "", "HTS 4819.10.00", ""),
        ("Section 301 duty (9903.05.36)", [f"=INDEX(eng_cv_{c},sel_row)*INDEX(eng_s301,sel_row)" for c in SKU_CODES], "", "USTR final action", ""),
        ("MPF (entry-level min/max, allocated by value)", [f"=INDEX(eng_MPF,sel_row)*INDEX(eng_cv_{c},sel_row)/INDEX(eng_CV,sel_row)" for c in SKU_CODES], "", "19 CFR 24.23; CBP Dec. 26-14", ""),
        ("HMF 0.125%", [f"=INDEX(eng_HMF,sel_row)*INDEX(eng_cv_{c},sel_row)/INDEX(eng_CV,sel_row)" for c in SKU_CODES], "", "19 CFR 24.24", ""),
        ("Duty + MPF + HMF per unit shipped", [f"=IF(INDEX(eng_units_{c},sel_row)>0,(INDEX(eng_cv_{c},sel_row)*INDEX(eng_duty_rate,sel_row)+(INDEX(eng_MPF,sel_row)+INDEX(eng_HMF,sel_row))*INDEX(eng_cv_{c},sel_row)/INDEX(eng_CV,sel_row))/INDEX(eng_units_{c},sel_row),0)" for c in SKU_CODES], "", "", "per box"),
    ]
    for lab, fs, trt, auth, note in lines:
        put(ws, r, 2, lab, wrap=True, bold=lab.startswith("CUSTOMS"))
        for j, f in enumerate(fs):
            calc_cell(ws, r, 3 + j, f, FMT_USD4 if "per unit" in lab else FMT_USD, bold=lab.startswith("CUSTOMS"))
        calc_cell(ws, r, 6, f"=SUM(C{r}:E{r})" if "per unit" not in lab else '=""', FMT_USD)
        put(ws, r, 8, trt)
        put(ws, r, 9, auth, wrap=True)
        put(ws, r, 10, note, wrap=True)
        if lab.startswith("CUSTOMS"):
            book.name("tar_cv_total", ws.title, f"F{r}")
            calc_cell(ws, r, 7, f"=IF(ABS(F{r}-SUM(F{r - 6}:F{r - 1}))<0.01,\"OK\",\"FAIL\")")
            book.name("tar_cv_check", ws.title, f"G{r}")
        r += 1
    r += 1
    r = section(ws, r, "5. Duty alternatives – per box effect at selected case (formula)", 10)
    r = header(ws, r, ["", "Treatment", "Rate", "Duty USD (shipment)", "Duty per saleable box", "Δ GM vs model (pts)", "", "Status", "", "Evidence needed"])
    for lab, rate, status, need in [("MFN only (Section 301 exempt or invalidated)", 0.0, "Legally supported alternative", "Broker Annex II check; CIT outcome"),
                                    ("Section 301 at 12.5% (model BASE)", 0.125, "Most likely applicable", "Broker confirmation"),
                                    ("Hypothetical escalation to 25%", 0.25, "Not in force – stress test only", "n/a")]:
        put(ws, r, 2, lab, wrap=True)
        inp_cell(ws, r, 3, rate, FMT_PCT)
        calc_cell(ws, r, 4, f"=INDEX(eng_CV,sel_row)*(INDEX(in_cus_mfn,1,sel_scen)+C{r})", FMT_USD)
        calc_cell(ws, r, 5, f"=D{r}/INDEX(eng_SALE,sel_row)", FMT_USD4)
        calc_cell(ws, r, 6, f"=(INDEX(eng_DUTY,sel_row)-D{r})/INDEX(eng_REV,sel_row)", "+0.0%;-0.0%")
        put(ws, r, 8, status, wrap=True)
        put(ws, r, 10, need, wrap=True)
        r += 1


def build_competitors(book):
    ws = book.ws(SH["comp"])
    title(ws, "Competitor Benchmark – US prices for comparable pizza boxes",
          "Advertised prices are NOT negotiated transaction prices. Comparability 1–5 (5 = same spec). Achievable wholesale price is derived by formula from the inputs in section 2.")
    set_widths(ws, {"A": 22, "B": 38, "C": 14, "D": 22, "E": 12, "F": 7, "G": 11, "H": 24, "I": 10, "J": 16, "K": 10, "L": 22, "M": 14, "N": 12, "O": 7, "P": 8, "Q": 50})
    r = section(ws, 4, "1. Listings collected (search extracts; access 2026-10-16)", 17)
    r = header(ws, r, ["Supplier", "Product", "Dimensions", "Board", "Printing", "Pack", "MOQ", "Published price", "$/box", "Shipping incl.?", "Tax incl.?",
                       "Price type", "Availability", "Collected", "Comparability", "Source", "Notes / URL"])
    for c in R.COMPETITORS:
        vals = [c["supplier"], c["product"], c["dims"], c["board"], c["printing"], c["per_pack"], c["moq"], c["price"], c["ppb"], c["shipping"], c["tax"],
                c["ptype"], c["avail"], c["date"], c["comparability"], c["sid"], (c["note"] + " " + c["url"]).strip()]
        for j, v in enumerate(vals):
            if j == 8 and v is not None:
                static_cell(ws, r, 1 + j, v, FMT_USD4)
            else:
                put(ws, r, 1 + j, v if v is not None else "n/a", wrap=True)
        ws.row_dimensions[r].height = 42
        r += 1
    r += 1
    r = section(ws, r, "2. Market inputs (scenario-dependent)", 10)
    items = [(f"listing_{k}", I.MARKET["listing"][k]) for k in SKU_CODES] + [("eflute_disc", I.MARKET["eflute_disc"]), ("dist_margin", I.MARKET["dist_margin"])]
    rows, r = input_table(book, ws, r, "mkt", items, "MKT")
    r += 1
    r = section(ws, r, "3. Achievable wholesale price per box (what a container-lot buyer could pay) = listing × (1 − E-flute discount if listing is B-flute) × (1 − buyer resale margin)", 10)
    r = header(ws, r, ["SKU", "Benchmark listing flute", "", "LOW", "BASE", "STRESS", "ACTIVE", "", "", "Note"])
    for code in SKU_CODES:
        put(ws, r, 1, code, bold=True)
        put(ws, r, 2, f"=sku_flute_{code}")
        for j in range(3):
            sc = j + 1
            calc_cell(ws, r, 4 + j, f"=INDEX(in_mkt_listing_{code},1,{sc})*(1-IF(sku_flute_{code}=\"B\",INDEX(in_mkt_eflute_disc,1,{sc}),0))*(1-INDEX(in_mkt_dist_margin,1,{sc}))", FMT_USD4)
        calc_cell(ws, r, 7, f"=INDEX(D{r}:F{r},1,sel_scen)", FMT_USD4, bold=True)
        book.name(f"mkt_price_{code}", ws.title, f"D{r}:F{r}")
        r += 1
    r += 1
    r = section(ws, r, "4. Import-parity cross-check (petition data: subject imports landed duty-paid $68.3m / 92.4m lb in 2025)", 10)
    r = header(ws, r, ["Measure", "", "", "3464", "3463", "3460", "", "", "", "Note"])
    put(ws, r, 1, "Subject-import LDP $/lb (2025)")
    inp_cell(ws, r, 4, 68.3 / 92.4, "0.000")
    book.name("mkt_ldp_lb", ws.title, f"D{r}")
    put(ws, r, 10, "S099 – petitioner data; mix of flutes/prints unknown", wrap=True)
    r += 1
    put(ws, r, 1, "Import-parity value per BigBoxPizza box (active weight)")
    for j, code in enumerate(SKU_CODES):
        calc_cell(ws, r, 4 + j, f"=mkt_ldp_lb*INDEX(lg_kpu_{code},sel_scen)/0.45359237", FMT_USD4)
    put(ws, r, 10, "Per-lb parity × estimated box weight (landed at port, before importer margin and delivery)", wrap=True)
    r += 1
    put(ws, r, 1, "BigBoxPizza cleared-port cost per lb (selected case)")
    calc_cell(ws, r, 4, "=(INDEX(eng_COGS,sel_row)-INDEX(eng_DELIV,sel_row)-INDEX(eng_UNLOAD,sel_row))/(INDEX(eng_kg,sel_row)/0.45359237)", "0.000")
    put(ws, r, 10, "Compare with subject-import LDP $/lb above", wrap=True)
    r += 2
    r = section(ws, r, "5. Competitive landscape notes", 10)
    for t in ["US production ≈ 3 billion pizza boxes a year (Pratt, S098); Smurfit Westrock and Pratt Industries (petitioners) are among the largest US integrated box makers.",
              "Imports from China, Malaysia and Türkiye rose ~130% (2023→2025) and face AD/CVD petitions (allegedly dumped 97–569%). Egypt is not named; a surge of Egyptian boxes could attract a future petition.",
              "US standard for 14\"–18\" is B-flute (~1/8\"); E-flute (~1/16\") is mainstream for 10\"–12\" and niche for 14\"/16\" (Arvco via Imperial Dade). NFC's 'SINGLE-E' 14\"/16\" must be validated for rigidity and heat retention.",
              "Realistic buyers: pizza-specialist distributors (e.g., PFG's Roma in Elizabeth, NJ – S085), broadline/disposables distributors, cash-and-carry, regional multi-unit pizza groups and commissaries. No demand or commitments were obtained."]:
        put(ws, r, 1, t, wrap=True)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=17)
        ws.row_dimensions[r].height = 30
        r += 1


def build_sku_master(book):
    ws = book.ws(SH["sku"])
    title(ws, "Product & SKU Master", "Descriptions preserved exactly as printed. Dimensions from the quotation (mm, inside/outside UNCONFIRMED).")
    set_widths(ws, {"A": 7, "B": 24, "C": 26, "D": 24, "E": 50, "F": 7, "G": 7, "H": 7, "I": 8, "J": 8, "K": 8, "L": 9, "M": 10, "N": 10, "O": 12, "P": 12, "Q": 10, "R": 10, "S": 11, "T": 12, "U": 10})
    r = header(ws, 4, ["SKU", "Description (verbatim Arabic)", "Literal English rendering", "Brief's transcription", "Discrepancy / flag", "L mm", "W mm", "H mm",
                       "L in", "W in", "H in", "US nominal", "Plan area m²", "Unit EGP", "EGP per m² plan", "USD per m² plan", "Est. blank L (BASE)",
                       "Est. blank W", "Est. board m² (net)", "EGP per m² board", "Benchmark flute"])
    for k in I.SKUS:
        code = k["code"]
        put(ws, r, 1, code, bold=True)
        put(ws, r, 2, k["desc_ar"], wrap=True)
        put(ws, r, 3, k["desc_doc_en"], wrap=True)
        put(ws, r, 4, k["desc_brief"], wrap=True)
        put(ws, r, 5, k["flag"], wrap=True, font=F_WARN)
        calc_cell(ws, r, 6, f"=q_L_{code}", "0")
        calc_cell(ws, r, 7, f"=q_W_{code}", "0")
        calc_cell(ws, r, 8, f"=q_H_{code}", "0")
        calc_cell(ws, r, 9, f"=F{r}/25.4", "0.00")
        calc_cell(ws, r, 10, f"=G{r}/25.4", "0.00")
        calc_cell(ws, r, 11, f"=H{r}/25.4", "0.00")
        put(ws, r, 12, k["us_nominal"])
        calc_cell(ws, r, 13, f"=F{r}*G{r}/1000000", "0.0000")
        calc_cell(ws, r, 14, f"=q_price_{code}", "#,##0.000")
        calc_cell(ws, r, 15, f"=N{r}/M{r}", FMT_NUM)
        calc_cell(ws, r, 16, f"=O{r}/fx_active", FMT_USD)
        calc_cell(ws, r, 17, f"=INDEX(lg_blankL_{code},2)", "0")
        calc_cell(ws, r, 18, f"=INDEX(lg_blankW_{code},2)", "0")
        calc_cell(ws, r, 19, f"=Q{r}*R{r}/1000000*INDEX(in_eng_net_area,1,2)", "0.0000")
        calc_cell(ws, r, 20, f"=N{r}/S{r}", FMT_NUM)
        inp_cell(ws, r, 21, k["bench_flute"])
        book.name(f"sku_flute_{code}", ws.title, f"U{r}")
        ws.row_dimensions[r].height = 48
        r += 1
    list_validation(ws, f"U5:U{r - 1}", ["B", "E"])
    r += 1
    put(ws, r, 1, "Notes", bold=True)
    for t in ["Plan-area prices are a preliminary comparison only; they do not measure board consumption. Board area uses the PROVISIONAL blank estimate (Loading Geometry).",
              "305/355/405 mm ≈ 12.0/14.0/15.9 in. US pizza boxes are specified by INSIDE size; if these are outside dimensions the usable size is smaller.",
              "Benchmark flute = flute of the US listing used as price benchmark (B-flute listings are discounted to E-flute equivalence)."]:
        r += 1
        put(ws, r, 2, t, wrap=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=16)
        ws.row_dimensions[r].height = 28


def build_spec(book):
    ws = book.ws(SH["spec"])
    title(ws, "Technical Specifications – what is known, what is assumed, what must be tested",
          "'SINGLE-E' most likely means single-wall E-flute corrugated board, but it does not prove construction, grammage, food-contact status or performance.")
    set_widths(ws, {"A": 9, "B": 44, "C": 18, "D": 18, "E": 18, "F": 14, "G": 26, "H": 18, "I": 16, "J": 50})
    r = section(ws, 4, "1. Specification status per SKU", 10)
    r = header(ws, r, ["#", "Specification item", "3464", "3463", "3460", "Status", "Provisional value in model", "Evidence", "Who supplies", "Why it matters"])
    items = [
        ("Finished internal dimensions", "unknown", "UNRESOLVED", "Quoted L×W×H used", "Fit of 12/14/16\" pizzas"),
        ("Finished external dimensions", "355×355×50 / 405×405×50 / 305×305×50 (basis unknown)", "UNRESOLVED", "As quoted", "Stacking, shelf space"),
        ("Lid height / base depth", "50 mm total (split unknown)", "UNRESOLVED", "50 mm", "Topping clearance"),
        ("Die-cut blank dimensions / flat pattern", "not supplied", "UNRESOLVED", "Estimated 2L+4.5H × W+2H+13", "Container capacity"),
        ("Board type / flute profile", "'SINGLE-E' (printed)", "UNRESOLVED", "Single-wall E-flute", "Rigidity, insulation"),
        ("Liner & medium composition, grammage", "not supplied", "UNRESOLVED", "≈420 g/m² combined (calibrated)", "Weight, strength, cost"),
        ("Recycled vs virgin fibre; food-contact surface", "not supplied", "UNRESOLVED", "Recycled likely", "21 CFR 176.260 / buyer acceptance"),
        ("Total caliper (thickness)", "not supplied", "UNRESOLVED", "1.6 mm stacking pitch", "Units per container"),
        ("Moisture / grease / wet-strength performance", "not supplied", "UNRESOLVED", "Not assumed", "Hot, greasy pizza"),
        ("Edge crush / compression strength", "not supplied", "UNRESOLVED", "Not assumed", "Stacking in delivery"),
        ("Printing, inks, adhesives, coatings", "not supplied (likely plain)", "UNRESOLVED", "Plain", "Food contact, TPCH, odour"),
        ("Tolerances", "not supplied", "UNRESOLVED", "–", "Fit, machinability"),
        ("Food-contact declarations & test data", "not supplied", "UNRESOLVED", "Testing budget in selling costs", "Legal & buyer requirement"),
        ("PFAS-free declaration / total fluorine", "not supplied", "UNRESOLVED", "Test in budget", "NY ECL 37-0209, MD law"),
        ("Heavy metals (Pb+Cd+Hg+Cr⁶⁺ ≤100 ppm)", "not supplied", "UNRESOLVED", "Test in budget", "TPCH states incl. NY/NJ"),
        ("Blanks per bundle; bundle dims & weight", "not supplied", "UNRESOLVED", "50; estimated", "Loading, handling"),
        ("Palletisation", "not specified", "UNRESOLVED", "Floor-loaded bundles", "Unloading, capacity"),
    ]
    for i, (it, val, st, prov, why) in enumerate(items, 1):
        put(ws, r, 1, f"T{i:02d}")
        put(ws, r, 2, it, wrap=True)
        for j in range(3):
            put(ws, r, 3 + j, val, wrap=True)
        put(ws, r, 6, st)
        put(ws, r, 7, prov, wrap=True)
        put(ws, r, 8, "UNRESOLVED")
        put(ws, r, 9, "NFC / lab")
        put(ws, r, 10, why, wrap=True)
        ws.row_dimensions[r].height = 30
        r += 1
    evidence_cf(ws, f"F{r - len(items)}:F{r - 1}")
    r += 1
    r = section(ws, r, "2. Blank & bundle estimation parameters (used by Loading Geometry)", 10)
    rows, r = input_table(book, ws, r, "eng", list(I.BLANK.items()), "ENG")
    r += 1
    r = section(ws, r, "3. Specification-equivalence matrix (US comparables)", 10)
    r = header(ws, r, ["#", "Comparable product", "Size", "Flute", "Fibre", "Weight / box", "Pack", "Price $/box", "Reliable comparison?", "Comment"])
    for i, row in enumerate([
        ("WebstaurantStore Choice 12×12×2 E-flute kraft", '12"', "E", "100% recycled", "≈136 g", "50", "$0.30–0.32", "YES for 3460", "Same flute; PFAS-free claim; online price"),
        ("WebstaurantStore Choice 12×12×2 B-flute kraft", '12"', "B", "Recycled", "n/a", "50", "$0.30–0.40", "Partial", "Heavier board; used to derive E/B ratio 0.80"),
        ("WebstaurantStore Choice 14×14×2 B-flute kraft", '14"', "B", "Recycled", "≈189 g", "50", "$0.39–0.42", "Partial for 3464", "Stronger board than SINGLE-E – discounted 20%"),
        ("WebstaurantStore Choice 16×16×2 kraft/white", '16"', "B (class)", "n/a", "n/a", "50", "$0.45–0.46", "Partial for 3463", "Discounted 20% to E-flute equivalence"),
        ("Arvco 14\"/16\" E-flute kraft (Imperial Dade)", '14"/16"', "E", "n/a", "n/a", "50", "login", "YES (spec) – no price", "Shows E-flute 14/16 is sold in US"),
        ("Uline 12\"/14\" white", '12"/14"', "B (white)", "n/a", "n/a", "50", "$0.78–1.08", "NO", "Premium white industrial catalogue"),
    ], 1):
        put(ws, r, 1, f"E{i:02d}")
        for j, v in enumerate(row):
            put(ws, r, 2 + j, v, wrap=True)
        r += 1
    r += 1
    r = section(ws, r, "4. Minimum pre-shipment testing plan (commonly used methods – not legally mandatory unless stated)", 10)
    r = header(ws, r, ["#", "Test", "Method (indicative)", "Why", "", "Mandatory?", "Indicative cost", "Evidence", "Source", "Note"])
    for i, row in enumerate([
        ("Dimensions & fit with real 12/14/16\" pizzas", "Measure inside dims; trial fit", "Size risk (inside vs outside)", "No – commercial", "$0–200", "ASSUMED", "", "Do on samples first"),
        ("Board construction, grammage, caliper", "ISO 536 (grammage), ISO 3034 (thickness)", "Confirms SINGLE-E spec & weight", "No", "$100–300", "ASSUMED", "", "Analyst knowledge – verify with lab"),
        ("Edge crush / flat crush", "ISO 3037 / TAPPI T 811 (ECT); ISO 3035 (FCT)", "Strength vs US B-flute norm", "No – buyer spec", "$100–300", "ASSUMED", "", ""),
        ("Water absorption (Cobb) & grease resistance", "ISO 535 / TAPPI T 441 (Cobb); TAPPI T 559 (kit test)", "Hot/greasy pizza performance", "No", "$100–300", "ASSUMED", "", "Only if a grease-resistance claim is made"),
        ("PFAS screening", "Total (organic) fluorine screening", "NY/MD PFAS bans; buyer requirement", "Legal in NY/MD (no intentionally added PFAS)", "a few hundred $ per sample", "ESTIMATED", "S092", ""),
        ("Heavy metals (Pb, Cd, Hg, Cr⁶⁺)", "Total metals by acid digestion + ICP", "TPCH ≤100 ppm combined", "Legal in TPCH states", "$150–300", "ASSUMED", "S035", ""),
        ("Food-contact documentation review", "Supplier DoC vs 21 CFR 176.170 / 176.260 / 175.105", "FDA food-additive compliance", "Compliance with FDA regs is legally required; the format of evidence is commercial", "$0–1,000 (consultant)", "ASSUMED", "S030", ""),
        ("Odour / taint (if printed or recycled)", "EN 1230 sensory test", "Food acceptability", "No", "$200–500", "ASSUMED", "", ""),
        ("Assembly & flat-pack integrity; real pizza trial", "Practical trial at a pizzeria", "Usability & heat/steam behaviour", "No", "$0–200", "ASSUMED", "", "Do with target buyer"),
    ], 1):
        put(ws, r, 1, f"Q{i:02d}")
        for j, v in enumerate(row):
            put(ws, r, 2 + j if j < 3 else 3 + j, v, wrap=True)
        ws.row_dimensions[r].height = 34
        r += 1


def build_engine(book):
    """Calc Engine: one row per scenario × configuration × destination, plus sensitivity and combo rows."""
    ws = book.ws(SH["eng"])
    title(ws, "Calc Engine – all scenario × configuration × destination cases (formula-driven)",
          "Each row recomputes the full landed-cost, margin and cash model from the input sheets. Driver columns (blue) override scenario inputs only on sensitivity/combination rows. Do not edit formulas.")
    cols = V.engine_columns()
    colmap = {}
    for j, (key, hdr, fmt, tmpl) in enumerate(cols, 1):
        colmap[key] = col(j)
    hdr_row = 4
    header(ws, hdr_row, [c[1] for c in cols])
    ws.freeze_panes = "F5"
    for j in range(1, len(cols) + 1):
        ws.column_dimensions[col(j)].width = 11
    ws.column_dimensions["B"].width = 34
    rows = []
    for si in range(1, 4):
        for ci in range(1, len(I.CONFIGS) + 1):
            for di in range(1, len(I.DESTINATIONS) + 1):
                rows.append(dict(key=si * 1000 + ci * 10 + di, label=f"{I.SCENARIOS[si - 1]} | {I.CONFIGS[ci - 1]['key']} | {I.DESTINATIONS[di - 1]['key']}",
                                 scen=si, cfg=ci, dest=di, ov={}))
    sens = E.sensitivity_cases(1)
    for i, (label, ov) in enumerate(sens, 1):
        xov = {}
        for k, v in ov.items():
            if k == "fx":
                which = 1 if "weaker" in label else 3
                xov[k] = f"INDEX(in_fx_cbe,1,{which})*(1-INDEX(in_fx_spread,1,{{c:scen}}))"
            elif k == "oceanfac":
                which = 1 if "−" in label else 3
                xov[k] = f"INDEX(in_oce_rate_hc,1,{which})/INDEX(in_oce_rate_hc,1,2)"
            else:
                xov[k] = v
        rows.append(dict(key=900000 + i, label="SENS: " + label, scen="sel_scen", cfg="sel_cfg", dest="sel_dest", ov=xov))
    for i, (label, ci, ov) in enumerate(E.combo_cases(), 1):
        xov = {}
        for k, v in ov.items():
            if k == "oceanfac":
                xov[k] = "INDEX(in_oce_rate_hc,1,1)/INDEX(in_oce_rate_hc,1,2)"
            else:
                xov[k] = v
        rows.append(dict(key=950000 + i, label="COMBO: " + label, scen=2, cfg=ci + 1, dest="sel_dest", ov=xov))
    r0 = hdr_row + 1
    driver_keys = set(E.DEFAULT_DRIVERS)
    for i, rw in enumerate(rows):
        r = r0 + i
        for j, (key, hdr, fmt, tmpl) in enumerate(cols, 1):
            if key == "key":
                put(ws, r, j, rw["key"], "0")
                continue
            if key == "label":
                put(ws, r, j, rw["label"])
                continue
            if key in ("scen", "cfg", "dest"):
                v = rw[key]
                if isinstance(v, str):
                    calc_cell(ws, r, j, "=" + v, "0")
                else:
                    put(ws, r, j, v, "0")
                continue
            if key in rw["ov"]:
                v = rw["ov"][key]
                if isinstance(v, str):
                    f = "=" + expand(v, colmap, r, f"{colmap['scen']}{r}")
                    c = put(ws, r, j, f, fmt, font=F_INPUT, fill=FILL_INPUT)
                else:
                    c = put(ws, r, j, v, fmt, font=F_INPUT, fill=FILL_INPUT)
                continue
            f = "=" + expand(tmpl, colmap, r, f"{colmap['scen']}{r}")
            c = ws.cell(row=r, column=j, value=f)
            c.number_format = fmt
            c.font = F_NORM
            if key in driver_keys:
                c.fill = FILL_INPUT
    rl = r0 + len(rows) - 1
    for key, hdr, fmt, tmpl in cols:
        book.name(f"eng_{key}", ws.title, f"{colmap[key]}{r0}:{colmap[key]}{rl}")
    return dict(r0=r0, rl=rl, colmap=colmap, rows=rows)


def write_static_and_images():
    """Run the geometry module to produce static values and diagrams."""
    static = dict(dp={}, fullmix={}, images=[])
    for si in range(3):
        for k in I.SKUS:
            for ck in CT_KEYS:
                ct = LM.container_by_key(ck)
                static["dp"][(k["code"], ck, si)] = LM.capacity(k, ct, si)["dp_positions"]
        for i, c in enumerate(I.CONFIGS):
            if c["kind"] == "fullmix":
                static["fullmix"][(i, si)] = LM.best_equal_mix(LM.container_by_key(c["ctype"]), si)
    out = os.path.join(PROJ, "outputs")
    os.makedirs(out, exist_ok=True)
    q = {k["code"]: k["qty"] for k in I.SKUS}
    LM.draw_mixed(q, LM.container_by_key("40HC"), 1, "Quoted order 60,000 (20,000 each) – one 40' high-cube, BASE", os.path.join(out, "loading_quoted_40HC_BASE.png"))
    LM.draw_mixed(q, LM.container_by_key("40GP"), 1, "Quoted order 60,000 – one 40' GP, BASE", os.path.join(out, "loading_quoted_40GP_BASE.png"))
    LM.draw_mixed(q, LM.container_by_key("40GP"), 2, "Quoted order 60,000 – one 40' GP, STRESS (1.8 mm pitch)", os.path.join(out, "loading_quoted_40GP_STRESS.png"))
    t = static["fullmix"][(3, 1)]
    LM.draw_mixed({k["code"]: t for k in I.SKUS}, LM.container_by_key("40HC"), 1, f"Optimised equal mix – {3 * t:,} boxes in one 40' HC, BASE", os.path.join(out, "loading_fullmix_40HC_BASE.png"))
    static["images"] = ["loading_quoted_40HC_BASE.png", "loading_quoted_40GP_BASE.png", "loading_quoted_40GP_STRESS.png", "loading_fullmix_40HC_BASE.png"]
    return static


def main(path):
    static = write_static_and_images()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for k in ORDER:
        wb.create_sheet(SH[k])
    book = Book(wb)
    V.SH = SH
    V.build_dashboard(book)          # creates sel_* names first
    V.build_readme(book)
    build_sku_master(book)
    build_quote(book)
    build_spec(book)
    build_fx(book)
    build_containers(book)
    build_geometry(book, static)
    build_shipment(book, static)
    build_origin(book)
    build_ocean(book)
    build_us_customs(book)
    build_destination(book)
    build_tariffs(book)
    eng = build_engine(book)
    V.build_landed(book)
    V.build_pricing(book)
    V.build_profitability(book)
    V.build_sensitivity(book, eng)
    V.build_scenarios(book)
    V.build_cash(book)
    build_competitors(book)
    V.build_risk(book)
    V.build_assumptions(book)
    V.build_sources(book)
    V.build_rfq(book)
    V.build_checks(book, eng)
    wb.calculation.fullCalcOnLoad = True
    for wsx in wb.worksheets:
        wsx.sheet_properties.tabColor = {"Executive Dashboard": "1F3864", "Calc Engine": "7F7F7F", "Model Checks": "C00000"}.get(wsx.title, "8EA9DB")
    wb.save(path)
    return path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(PROJ, "outputs", "BigBoxPizza_US_Landed_Cost_Model_formulas.xlsx")
    print("written", main(out))
