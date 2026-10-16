"""Recalculate the workbook in LibreOffice (headless, via UNO) and validate it.

Usage: python3 validate_model.py <formula_workbook.xlsx> <recalculated_output.xlsx> [report.md]

Checks
 1. Recalculation succeeds; every sheet re-opens.
 2. No error values (#VALUE!, #REF!, #NAME?, #DIV/0!, #N/A, Err:nnn) anywhere.
 3. Every Calc Engine row equals the independent Python engine (engine.py) on key outputs.
 4. Loading-geometry capacities and configuration feasibility equal loading_model.py.
 5. Model Checks sheet: list of statuses; overall status.
 6. Required worksheet names present; formula counts per sheet.
"""
import json
import math
import os
import socket
import subprocess
import sys
import time

import openpyxl

import engine as E
import inputs as I
import loading_model as LM

REQUIRED = ["Executive Dashboard", "Read Me & Instructions", "Product & SKU Master", "Supplier Quotation", "Technical Specifications", "Exchange Rates",
            "Container Specifications", "Loading Geometry", "Shipment Configuration", "Origin Costs", "Ocean Freight", "U.S. Port & Customs",
            "Destination Delivery", "Tariffs & Customs Value", "Landed Cost by SKU", "Pricing & Margin", "Container Profitability", "Sensitivity Analysis",
            "Scenario Comparison", "Cash Flow & Working Capital", "Competitor Benchmark", "Risk Register", "Assumptions Register", "Source Register",
            "Quote Requests & Missing Data"]


def free_port():
    s = socket.socket()
    s.bind(("localhost", 0))
    p = s.getsockname()[1]
    s.close()
    return p


def recalc(src, dst):
    import uno
    from com.sun.star.beans import PropertyValue

    def prop(n, v):
        p = PropertyValue()
        p.Name = n
        p.Value = v
        return p

    port = free_port()
    profile = "file:///tmp/lo_profile_bbp_" + str(port)
    proc = subprocess.Popen(["soffice", "--headless", "--invisible", "--nologo", "--norestore", "-env:UserInstallation=" + profile,
                             f"--accept=socket,host=localhost,port={port};urp;"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    ctx = None
    for _ in range(90):
        try:
            local = uno.getComponentContext()
            resolver = local.ServiceManager.createInstanceWithContext("com.sun.star.bridge.UnoUrlResolver", local)
            ctx = resolver.resolve(f"uno:socket,host=localhost,port={port};urp;StarOffice.ComponentContext")
            break
        except Exception:
            time.sleep(1)
    if ctx is None:
        proc.kill()
        raise RuntimeError("LibreOffice did not start")
    desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
    doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(os.path.abspath(src)), "_blank", 0, (prop("Hidden", True),))
    doc.calculateAll()
    doc.calculateAll()
    doc.storeToURL(uno.systemPathToFileUrl(os.path.abspath(dst)), (prop("FilterName", "Calc MS Excel 2007 XML"),))
    doc.close(True)
    try:
        desktop.terminate()
    except Exception:
        pass
    try:
        proc.wait(timeout=60)
    except Exception:
        proc.kill()


def lo_version():
    try:
        out = subprocess.run(["soffice", "--version"], capture_output=True, text=True, timeout=60).stdout.split()
        return " ".join(out[:2]) if out else "LibreOffice (version unknown)"
    except (OSError, subprocess.SubprocessError):
        return "LibreOffice (version unknown)"


def close(a, b, tol=1e-6):
    if a is None and b in (None, "", "n/a"):
        return True
    if isinstance(b, str) and (a is None or a == "n/a"):
        return b in ("", "n/a")
    if isinstance(a, bool) or isinstance(b, bool):
        return bool(a) == bool(b)
    try:
        a, b = float(a), float(b)
    except (TypeError, ValueError):
        return a == b
    return abs(a - b) <= tol * max(1.0, abs(a), abs(b))


def main(src, dst, report):
    t0 = time.time()
    recalc(src, dst)
    res = dict(recalc_seconds=round(time.time() - t0, 1), source=os.path.basename(src), recalculated=os.path.basename(dst), libreoffice=lo_version())
    wbv = openpyxl.load_workbook(dst, data_only=True)
    wbf = openpyxl.load_workbook(src)
    # ---- 6. sheets and formula counts
    res["sheets"] = wbv.sheetnames
    res["missing_required_sheets"] = [s for s in REQUIRED if s not in wbv.sheetnames]
    fcount = {}
    for ws in wbf.worksheets:
        n = 0
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n += 1
        fcount[ws.title] = n
    res["formula_cells"] = fcount
    res["formula_cells_total"] = sum(fcount.values())
    # ---- 2. error scan
    errs = []
    for ws in wbv.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and (v.startswith("#") and v.rstrip("!?0/").upper() in ("#VALUE", "#REF", "#NAME", "#DIV", "#N/A", "#NUM", "#NULL") or v.startswith("Err:")
                                           or v in ("#VALUE!", "#REF!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")):
                    errs.append(f"{ws.title}!{c.coordinate}={v}")
    res["error_cells"] = errs[:50]
    res["error_count"] = len(errs)
    # ---- 3. engine comparison
    ws = wbv["Calc Engine"]
    hdr = {c.value: c.column for c in ws[4] if c.value}
    keys_hdr = {h: k for k, h, f, t in __import__("build_views").engine_columns()}
    col_of = {keys_hdr[h]: cidx for h, cidx in hdr.items() if h in keys_hdr}
    py = {}
    rows, base, sens, combos = E.run_all()
    for r in rows:
        py[(r["s"] + 1) * 1000 + (r["ci"] + 1) * 10 + (r["di"] + 1)] = r
    for i, (label, r) in enumerate(sens, 1):
        py[900000 + i] = r
    for i, (label, r) in enumerate(combos, 1):
        py[950000 + i] = r
    compare = ["U", "SUP", "VAT_CASH", "VAT_ECON", "ORIGIN", "OCEAN", "INS", "CV", "DUTY", "MPF", "HMF", "STB", "US_SHP", "US_CTR", "DELIV", "UNLOAD",
               "COGS", "REV", "GP", "GM", "SELL_VAR", "SELL_FIX", "FIN", "CONTRIB", "CM", "PEAK", "CASH_TOTAL", "avg_cpb", "avg_price", "f_be",
               "chg_gm_t", "chg_gm_0", "max_ocean", "u_be_t", "u_be_0"]
    per_sku = ["cpb", "price", "maxexw"]
    mism = []
    nrows = 0
    for rr in range(5, ws.max_row + 1):
        key = ws.cell(row=rr, column=1).value
        if key is None:
            continue
        key = int(key)
        if key not in py:
            mism.append(f"row {rr}: key {key} not in Python engine")
            continue
        pr = py[key]
        nrows += 1
        for k in compare:
            xv = ws.cell(row=rr, column=col_of[k]).value
            pv = pr[k]
            if not close(pv, xv, 1e-6):
                mism.append(f"key {key} {k}: excel={xv} python={pv}")
        for k in per_sku:
            for code in [s["code"] for s in I.SKUS]:
                xv = ws.cell(row=rr, column=col_of[f"{k}_{code}"]).value
                pv = pr[k][code]
                if not close(pv, xv, 1e-6):
                    mism.append(f"key {key} {k}_{code}: excel={xv} python={pv}")
        xf = ws.cell(row=rr, column=col_of["fits"]).value
        ci = pr["ci"]
        pf = LM.config_feasibility(I.CONFIGS[ci], pr["s"])["fits"]
        if bool(xf) != bool(pf):
            mism.append(f"key {key} fits: excel={xf} python={pf}")
    res["engine_rows_compared"] = nrows
    res["engine_values_per_row"] = len(compare) + len(per_sku) * len(I.SKUS) + 1  # + container-fit flag
    res["engine_mismatches"] = mism[:60]
    res["engine_mismatch_count"] = len(mism)
    # ---- 4. geometry
    wsg = wbv["Loading Geometry"]
    gm = []
    for rr in range(1, wsg.max_row + 1):
        a = wsg.cell(row=rr, column=1).value
        b = wsg.cell(row=rr, column=2).value
        d = wsg.cell(row=rr, column=4).value
        if a in [s["code"] for s in I.SKUS] and b in ("20GP", "40GP", "40HC") and isinstance(d, (int, float)):
            sku = [s for s in I.SKUS if s["code"] == a][0]
            cap = LM.capacity(sku, LM.container_by_key(b), int(d) - 1)
            xv = wsg.cell(row=rr, column=28).value
            if not close(cap["real_units"], xv):
                gm.append(f"{a} {b} s{d}: excel={xv} python={cap['real_units']}")
    res["geometry_mismatches"] = gm
    # ---- 5. model checks
    wsc = wbv["Model Checks"]
    checks = []
    for rr in range(5, wsc.max_row + 1):
        lab = wsc.cell(row=rr, column=2).value
        st = wsc.cell(row=rr, column=3).value
        if lab:
            checks.append((lab, st))
    res["model_checks"] = checks
    # ---- dashboard snapshot
    wsd = wbv["Executive Dashboard"]
    snap = {}
    for rr in range(1, 45):
        lab = wsd.cell(row=rr, column=1).value
        if lab:
            snap[str(lab)] = wsd.cell(row=rr, column=2).value
    res["dashboard_snapshot"] = {k: (round(v, 6) if isinstance(v, float) else (str(v) if v is not None else None)) for k, v in snap.items()}
    ok = (not res["missing_required_sheets"] and res["error_count"] == 0 and res["engine_mismatch_count"] == 0 and not gm)
    res["status"] = "PASS" if ok else "FAIL"
    with open(report.replace(".md", ".json"), "w", encoding="utf-8") as f:
        json.dump(res, f, indent=1, ensure_ascii=False, default=str)
    with open(report, "w", encoding="utf-8") as f:
        f.write(f"# Workbook validation report\n\nStatus: **{res['status']}**\n\n")
        f.write(f"- Source workbook: `{res['source']}`; recalculated with LibreOffice (headless, UNO calculateAll) in {res['recalc_seconds']} s → `{res['recalculated']}`\n")
        f.write(f"- Worksheets: {len(res['sheets'])} ({', '.join(res['sheets'])})\n")
        f.write(f"- Missing required worksheets: {res['missing_required_sheets'] or 'none'}\n")
        f.write(f"- Formula cells: {res['formula_cells_total']:,}\n")
        f.write(f"- Error values after recalculation: {res['error_count']}\n")
        f.write(f"- Calc Engine rows compared with independent Python engine: {res['engine_rows_compared']} (mismatches: {res['engine_mismatch_count']})\n")
        f.write(f"- Loading-geometry capacity mismatches: {len(gm)}\n\n")
        f.write("## Model Checks sheet\n\n| Check | Status |\n|---|---|\n")
        for lab, st in checks:
            f.write(f"| {lab} | {st} |\n")
        f.write("\n## Formula cells per sheet\n\n| Sheet | Formula cells |\n|---|---:|\n")
        for k, v in fcount.items():
            f.write(f"| {k} | {v:,} |\n")
        if mism or errs or gm:
            f.write("\n## Problems\n\n")
            for x in (errs[:30] + mism[:30] + gm):
                f.write(f"- {x}\n")
    return res


if __name__ == "__main__":
    src, dst = sys.argv[1], sys.argv[2]
    rep = sys.argv[3] if len(sys.argv) > 3 else os.path.join(os.path.dirname(os.path.abspath(dst)), "validation_report.md")
    r = main(src, dst, rep)
    print("STATUS", r["status"], "| errors", r["error_count"], "| engine rows", r["engine_rows_compared"], "mismatches", r["engine_mismatch_count"],
          "| geometry mismatches", len(r["geometry_mismatches"]), "| formulas", r["formula_cells_total"])
    for e in r["error_cells"][:25]:
        print("  ERR", e)
    for m in r["engine_mismatches"][:25]:
        print("  MISM", m)
    for lab, st in r["model_checks"]:
        print("  CHECK", st, "|", lab)
