"""Small openpyxl toolkit: styles, defined names, input tables, formula templates."""
import re

from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# ----------------------------------------------------------------------------- styles
NAVY = "1F3864"
F_TITLE = Font(name="Calibri", size=15, bold=True, color=NAVY)
F_SUB = Font(name="Calibri", size=10, italic=True, color="595959")
F_HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_BOLD = Font(name="Calibri", size=10, bold=True)
F_NORM = Font(name="Calibri", size=10)
F_INPUT = Font(name="Calibri", size=10, color="0000FF")
F_LINK = Font(name="Calibri", size=10, color="006100")
F_WARN = Font(name="Calibri", size=10, bold=True, color="9C0006")
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_SEC = PatternFill("solid", fgColor="D9E1F2")
FILL_INPUT = PatternFill("solid", fgColor="FFF2CC")
FILL_STATIC = PatternFill("solid", fgColor="FCE4D6")
FILL_CALC = PatternFill("solid", fgColor="F2F2F2")
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_BAD = PatternFill("solid", fgColor="FFC7CE")
FILL_AMBER = PatternFill("solid", fgColor="FFEB9C")
FILL_KPI = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

FMT_USD = '"$"#,##0.00'
FMT_USD0 = '"$"#,##0'
FMT_USD4 = '"$"#,##0.0000'
FMT_EGP = '#,##0.00 "EGP"'
FMT_EGP3 = '#,##0.000 "EGP"'
FMT_PCT = "0.0%"
FMT_PCT2 = "0.00%"
FMT_INT = "#,##0"
FMT_NUM = "#,##0.00"
FMT_NUM3 = "#,##0.000"
FMT_NUM4 = "#,##0.0000"


def fmt_for_unit(unit):
    u = unit.lower()
    if u.startswith("%") or u in ("% rev", "% units", "% p.a.") or u.startswith("%"):
        return FMT_PCT2
    if "egp/usd" in u:
        return FMT_NUM4
    if u.startswith("egp"):
        return '#,##0'
    if u.startswith("usd per $1,000"):
        return FMT_NUM
    if u.startswith("usd"):
        return FMT_USD
    if u in ("ratio", "×h"):
        return FMT_NUM3
    if u == "flag" or u == "mode":
        return "0"
    return "#,##0.###"


def q(sheet):
    return "'" + sheet.replace("'", "''") + "'"


def ref(sheet, cell_range):
    """Absolute reference string for defined names."""
    parts = cell_range.split(":")
    abs_parts = []
    for p in parts:
        m = re.match(r"([A-Z]+)(\d+)", p)
        abs_parts.append(f"${m.group(1)}${m.group(2)}")
    return f"{q(sheet)}!" + ":".join(abs_parts)


class Book:
    def __init__(self, wb):
        self.wb = wb
        self.names = {}

    def name(self, nm, sheet, cell_range):
        if nm in self.names:
            raise ValueError(f"duplicate defined name {nm}")
        self.names[nm] = (sheet, cell_range)
        self.wb.defined_names[nm] = DefinedName(nm, attr_text=ref(sheet, cell_range))

    def ws(self, title):
        return self.wb[title]


def title(ws, text, sub=None, width_cols=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB
        ws["A2"].alignment = Alignment(wrap_text=False)
    ws.sheet_view.showGridLines = False
    if width_cols:
        for col, w in width_cols.items():
            ws.column_dimensions[col].width = w


def section(ws, row, text, ncols=10):
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    for j in range(1, ncols + 1):
        ws.cell(row=row, column=j).fill = FILL_SEC
    return row + 1


def header(ws, row, labels, col0=1):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=col0 + j, value=lab)
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = CENTER
        c.border = BORDER
    return row + 1


def put(ws, row, col, value, fmt=None, font=None, fill=None, wrap=False, border=True, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    if fmt:
        c.number_format = fmt
    c.font = font or (F_BOLD if bold else F_NORM)
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = WRAP
    if border:
        c.border = BORDER
    return c


def inp_cell(ws, row, col, value, fmt=None):
    return put(ws, row, col, value, fmt=fmt, font=F_INPUT, fill=FILL_INPUT)


def static_cell(ws, row, col, value, fmt=None):
    return put(ws, row, col, value, fmt=fmt, font=F_NORM, fill=FILL_STATIC)


def calc_cell(ws, row, col, formula, fmt=None, bold=False):
    return put(ws, row, col, formula, fmt=fmt, font=F_BOLD if bold else F_NORM, fill=FILL_CALC)


def evidence_cf(ws, rng):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({rng.split(":")[0]},8)="VERIFIED"'], fill=FILL_OK))
    first = rng.split(":")[0]
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({first}="ESTIMATED",{first}="CALCULATED")'], fill=FILL_AMBER))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR({first}="ASSUMED",{first}="UNRESOLVED",{first}="UNVERIFIED")'], fill=FILL_BAD))


def status_cf(ws, rng):
    first = rng.split(":")[0]
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},2)="OK"'], fill=FILL_OK))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(LEFT({first},4)="FAIL",LEFT({first},5)="ERROR")'], fill=FILL_BAD))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},4)="WARN"'], fill=FILL_AMBER))


def list_validation(ws, cells, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(str(o) for o in options) + '"', allow_blank=False, showDropDown=False)
    dv.error = "Choose a value from the list"
    dv.errorTitle = "Invalid entry"
    ws.add_data_validation(dv)
    dv.add(cells)
    return dv


def number_validation(ws, cells, lo=None, hi=None):
    if lo is not None and hi is not None:
        dv = DataValidation(type="decimal", operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True)
    else:
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=str(lo if lo is not None else 0), allow_blank=True)
    dv.error = "Value outside the allowed range"
    ws.add_data_validation(dv)
    dv.add(cells)
    return dv


INPUT_HDR = ["ID", "Input", "Unit", "LOW", "BASE", "STRESS", "ACTIVE (selected scenario)", "Evidence", "Source IDs", "Notes / method"]


def input_table(book, ws, r0, prefix, items, id_prefix, sel_name="sel_scen"):
    """Write an input table and create names in_<prefix>_<key> (LOW:STRESS)."""
    r = header(ws, r0, INPUT_HDR)
    rows = {}
    for i, (key, d) in enumerate(items, 1):
        put(ws, r, 1, f"{id_prefix}-{i:02d}")
        put(ws, r, 2, d["name"], wrap=True)
        put(ws, r, 3, d["unit"])
        fmt = fmt_for_unit(d["unit"])
        for j, v in enumerate(d["v"]):
            inp_cell(ws, r, 4 + j, v, fmt)
        calc_cell(ws, r, 7, f"=INDEX(D{r}:F{r},1,{sel_name})", fmt)
        put(ws, r, 8, d["evidence"])
        put(ws, r, 9, d["sources"], wrap=True)
        put(ws, r, 10, d.get("note", ""), wrap=True)
        book.name(f"in_{prefix}_{key}", ws.title, f"D{r}:F{r}")
        rows[key] = r
        r += 1
    evidence_cf(ws, f"H{r0 + 1}:H{r - 1}")
    number_validation(ws, f"D{r0 + 1}:F{r - 1}", lo=-1e12, hi=1e12)
    return rows, r


def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# ----------------------------------------------------------------------------- formula templates
_TOK_C = re.compile(r"\{c:([A-Za-z0-9_]+)\}")
_TOK_IN = re.compile(r"\{in:([A-Za-z0-9_]+)\}")


def expand(template, colmap, row, scen_ref):
    def rc(m):
        key = m.group(1)
        if key not in colmap:
            raise KeyError(f"unknown engine column {key}")
        return f"{colmap[key]}{row}"

    def ri(m):
        return f"INDEX(in_{m.group(1)},1,{scen_ref})"

    out = _TOK_C.sub(rc, template)
    out = _TOK_IN.sub(ri, out)
    return out


def col(n):
    return get_column_letter(n)
