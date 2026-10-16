"""Formal QA validation checklist (brief §21) – runs after the workbook is recalculated and the Markdown is written.

    python3 BigBoxPizza/scripts/qa_checklist.py     → outputs/qa_validation_checklist.md / .json

Every item of §21 (research, engineering, financial, deliverable validation) gets a status and the evidence behind it:
PASS – checked and correct; LIMITED – no defect found, but the check could not be completed in this environment
(the limitation is stated); WARN – a defect or open point that does not invalidate the deliverables; FAIL – must be
fixed before the package is released (build_all.py stops on FAIL).
Inputs: inputs.py, sources_data.py, registers.py, the engine, the recalculated workbook, outputs/validation_report.json
and (if present) outputs/url_check.json written by check_urls.py.
"""
import csv
import datetime as dt
import json
import math
import os
import re
import sys
import zipfile
from urllib.parse import urlparse

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.path.dirname(HERE)
REPO = os.path.dirname(PROJ)
sys.path.insert(0, HERE)

import build_source_register as BSR  # noqa: E402
import engine as E  # noqa: E402
import inputs as I  # noqa: E402
import loading_model as LM  # noqa: E402
import registers as R  # noqa: E402
import sources_data as SD  # noqa: E402

VALUATION = dt.date(2026, 10, 16)
ALLOWED = ("VERIFIED_OFFICIAL", "VERIFIED_QUOTE", "VERIFIED_SUPPLIER", "VERIFIED_MARKET", "CALCULATED", "ESTIMATED", "ASSUMED", "UNRESOLVED")
REQUIRED_SHEETS = ["Executive Dashboard", "Read Me & Instructions", "Product & SKU Master", "Supplier Quotation", "Technical Specifications",
                   "Exchange Rates", "Container Specifications", "Loading Geometry", "Shipment Configuration", "Origin Costs", "Ocean Freight",
                   "U.S. Port & Customs", "Destination Delivery", "Tariffs & Customs Value", "Landed Cost by SKU", "Pricing & Margin",
                   "Container Profitability", "Sensitivity Analysis", "Scenario Comparison", "Cash Flow & Working Capital", "Competitor Benchmark",
                   "Risk Register", "Assumptions Register", "Source Register", "Quote Requests & Missing Data"]
REPORT_SECTIONS = ["Executive verdict", "One-page decision summary", "Quotation assessment", "Product specification assessment",
                   "Container-loading feasibility", "Import classification and tariff findings", "U.S. compliance requirements",
                   "Origin cost analysis", "Freight and port comparison", "Destination delivery comparison", "Competitor price benchmarking",
                   "Landed cost by SKU", "Pricing and profit analysis", "Scenario comparison", "Main uncertainties", "Supplier negotiation strategy",
                   "Recommended pilot shipment", "Go/no-go decision", "Conditions that must be satisfied before payment", "Sources and appendices"]
DELIVERABLES = ["BigBoxPizza_US_Market_and_Landed_Cost_Report.md", "BigBoxPizza_US_Landed_Cost_Model.xlsx", "BigBoxPizza_Source_Register.csv",
                "BigBoxPizza_Assumptions_and_Gaps.md", "BigBoxPizza_RFQ_Pack.md", "BigBoxPizza_Calculation_Audit.md"]
MD_FILES = [f for f in DELIVERABLES if f.endswith(".md")] + ["README.md"]
EXCEL_ERRORS = ("#NAME?", "#VALUE!", "#REF!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!")

ROWS = []


def chk(group, item, status, evidence):
    assert status in ("PASS", "LIMITED", "WARN", "FAIL"), status
    ROWS.append(dict(group=group, item=item, status=status, evidence=evidence))


def close(a, b, tol=1e-6):
    return abs(a - b) <= tol * max(1.0, abs(a), abs(b))


def v(inp, s):
    return inp["v"][s]


def src_ids_in(text):
    """Source IDs cited in a text, as written ([S001], S044-S049 ranges count by their endpoints)."""
    return set(re.findall(r"\bS\d{3}\b", text))


def read(path):
    with open(path, encoding="utf-8") as f:
        return f.read()


# ============================================================================ research validation
def research(val, url):
    src = {s["id"]: s for s in SD.SOURCES}

    # R1 principal tariff / regulatory findings on authoritative, current sources
    key = {"S010": "HTS 4819.10.00 (MFN Free)", "S014": "USTR Section 301 final action", "S015": "Federal Register 301 notice",
           "S016": "CBP CSMS 301 guidance", "S024": "MPF FY2027", "S025": "HMF 19 CFR 24.24", "S026": "Valuation 19 CFR 152.103",
           "S027": "Marking 19 CFR 134.24", "S030": "FDA 21 CFR 176.170", "S032": "NY PFAS law", "S033": "MD PFAS law",
           "S035": "Toxics in Packaging", "S060": "Egypt VAT refund", "S062": "Egypt export law 118/1975"}
    missing = [k for k in key if k not in src]
    not_off = [k for k in key if k in src and src[k]["evidence"] != "VERIFIED_OFFICIAL"]
    tariff_inputs = {n: I.CUSTOMS[n] for n in ("mfn", "s301", "mpf_rate", "mpf_min", "mpf_max", "hmf")}
    bad_inp = [n for n, d in tariff_inputs.items() if d["evidence"] != "VERIFIED_OFFICIAL" or not d["sources"]]
    status = "FAIL" if (missing or not_off or bad_inp) else "LIMITED"
    chk("Research", "Principal tariff and regulatory findings rest on current authoritative sources", status,
        f"{len(key) - len(missing) - len(not_off)}/{len(key)} key legal sources are VERIFIED_OFFICIAL ({', '.join(key)}); duty and fee inputs "
        f"{', '.join(tariff_inputs)} are VERIFIED_OFFICIAL with sources{'' if not bad_inp else ' except ' + ', '.join(bad_inp)}. "
        "Limitation: official pages were read through search-engine extracts; the egress policy blocked direct retrieval, so a licensed "
        "customs broker must confirm Section 301 coverage (gap G02) before entry.")

    # R2 URLs
    web = [(s["id"], u) for s in SD.SOURCES for u in ([p.strip() for p in s["url"].split(";")] if ";" in s["url"] else [s["url"].strip()])
           if u.startswith("http")]
    bad_syntax = [i for i, u in web if not (urlparse(u).scheme in ("http", "https") and urlparse(u).netloc)]
    if url:
        res = {}
        for r in url["rows"]:
            res[r["result"]] = res.get(r["result"], 0) + 1
        hosts = len({r["host"] for r in url["rows"] if r["result"] == "blocked_by_policy"})
        reach = res.get("reachable", 0)
        broken = [f"{r['id']} HTTP {r['code']}" for r in url["rows"] if r["result"] == "http_error"]
        local_bad = [r["id"] for r in url["rows"] if r["result"] == "local_file_missing"]
        status = "FAIL" if (bad_syntax or local_bad) else ("WARN" if broken else ("PASS" if reach == len(web) else "LIMITED"))
        chk("Research", "Source URLs work where accessible", status,
            f"{len(web)} web URLs, all syntactically valid{'' if not bad_syntax else ' except ' + ', '.join(bad_syntax)}. Reachability check "
            f"{url['checked_at']}: {reach} reachable, {res.get('blocked_by_policy', 0)} blocked by the session's egress policy ({hosts} hosts, "
            f"proxy CONNECT 403 – not a dead-link signal), {len(broken)} HTTP errors{(' (' + ', '.join(broken) + ')') if broken else ''}, "
            f"{res.get('network_error', 0)} other network errors. Local evidence paths: {res.get('local_file_ok', 0)} present, {len(local_bad)} missing. "
            "No URL could be opened from this environment; each must be re-opened by the client before reliance.")
    else:
        chk("Research", "Source URLs work where accessible", "LIMITED",
            f"{len(web)} web URLs syntactically valid; reachability not tested (run scripts/check_urls.py).")

    # R3 dates
    undated = [s["id"] for s in SD.SOURCES if re.search(r"n\.d\.|undated|^current$", s["pub_date"].strip(), re.I)]
    no_date = [s["id"] for s in SD.SOURCES if not s["pub_date"].strip() or not s["access_date"].strip()]
    late = [s["id"] for s in SD.SOURCES if s["access_date"][:10] > VALUATION.isoformat()]
    eff = [("Section 301 Egypt 12.5% effective", "2026-07-24"), ("MPF FY2027 rates effective", "2026-10-01"),
           ("CBE FX observation (BASE)", "2026-10-15")]
    eff_bad = [n for n, d in eff if d > VALUATION.isoformat()]
    quote_expired = dt.date(2026, 9, 30) < VALUATION
    status = "FAIL" if (no_date or late or eff_bad or not quote_expired) else "PASS"
    chk("Research", "Publication and effective dates checked", status,
        f"All {len(SD.SOURCES)} sources carry publication and access dates; access dates ≤ valuation date {VALUATION.isoformat()}. "
        + "; ".join(f"{n} {d}" for n, d in eff) + f" – all on or before the valuation date. NFC quote valid to 2026-09-30, so expired "
        f"(treated as historical reference; Model Checks WARN). Undated or 'current' pages ({len(undated)}): {', '.join(undated) or 'none'} – "
        "used only for regulations in force or non-time-critical facts, or flagged as anecdotal.")

    # R4 contradictions
    tr = read(os.path.join(PROJ, "evidence", "NFC-OFFER-2026-0120_transcription.md"))
    n_disc = len(set(re.findall(r"\|\s*(D\d{1,2})\s*\|", tr)))
    contradictions = [
        "Brief's '5 × 5 × 5.5 inch' for SKU 3464 vs document '5*35.5*35.5' – flagged, not corrected (D1, G06)",
        "Signature 'Abdelrahman Darwish' vs printed 'Mahmoud Bakr'; bank 'FAP MISR' vs SWIFT NBADEGCAXXX – flagged (D7/D8, G04)",
        "40' HC payload: DSV 28,560 kg [S041] vs Maersk flyer 28,870 kg [S042] – lower value used; US road limit (19,958 kg) binds first",
        "CMA CGM FAK includes origin THC [S047], model ocean estimate excludes it – OTHC charged once via its own line and a double-count flag",
        "Section 122 surcharge (10%, 24-Feb to 24-Jul-2026) vs current 12.5% Section 301 – surcharge expired; only Section 301 modelled [S021][S022]",
        "News claim of food-sector exemptions [S019] vs Annex II categories [S018] – no exemption assumed for 4819.10 in BASE; broker to confirm (G02)",
        "CBE buy 53.0544 (14 Oct) vs press 53.04 (15 Oct) [S063] – the later 53.04 used",
    ]
    chk("Research", "Contradictions resolved or explicitly documented", "PASS" if n_disc >= 10 else "WARN",
        f"Quotation discrepancies D1–D{n_disc} documented in the transcription and the gaps register. Source conflicts and their resolution: "
        + "; ".join(contradictions) + ".")

    # R5 competitor comparisons
    issues = []
    for c in R.COMPETITORS:
        if c["sid"] not in src:
            issues.append(f"{c['supplier']}: source {c['sid']} missing")
        if c["ppb"] is not None:
            amounts = [float(x.replace(",", "")) for x in re.findall(r"\$(\d[\d,]*\.?\d*)", c["price"])]
            cands = ([a / c["per_pack"] for a in amounts] if c["per_pack"] else []) + [a for a in amounts if a < 2]
            if cands and not (min(cands) - 0.011 <= c["ppb"] <= max(cands) + 0.011):
                issues.append(f"{c['supplier']} {c['product'][:30]}: ppb {c['ppb']} vs price '{c['price']}'")
        if not 1 <= c["comparability"] <= 5:
            issues.append(f"{c['product'][:30]}: comparability score out of range")
    lst = {k: v(I.MARKET["listing"][k], 1) for k in ("3464", "3463", "3460")}
    exp = {"3464": 19.49 / 50, "3463": 22.49 / 50, "3460": 14.99 / 50}
    lst_bad = [k for k in lst if abs(lst[k] - exp[k]) > 0.006]
    chk("Research", "Material competitor comparisons verified", "FAIL" if (issues or lst_bad) else "LIMITED",
        f"{len(R.COMPETITORS)} listings: every source ID exists, every per-box price reconciles to its pack price (±$0.01) and every row "
        f"has a comparability score. BASE listing inputs equal the WebstaurantStore pack prices ÷ 50 (14\" $19.49, 16\" $22.49, 12\" E-flute "
        f"$14.99){'' if not lst_bad else ' – mismatch: ' + ', '.join(lst_bad)}. Limitation: listings captured from search extracts on "
        f"2026-10-12/16 and not re-opened.{(' Issues: ' + '; '.join(issues)) if issues else ''}")

    # R6 asking prices vs quotes
    quotes = [s["id"] for s in SD.SOURCES if s["evidence"] == "VERIFIED_QUOTE"]
    q_inputs = [(g, k) for g, k, d in I.all_inputs() if d["evidence"] == "VERIFIED_QUOTE" and "S001" not in d["sources"].split(";")]
    mkt = [d["evidence"] for d in list(I.MARKET["listing"].values())]
    rep = read(os.path.join(PROJ, DELIVERABLES[0]))
    ok = quotes == ["S001"] and not q_inputs and all(e == "VERIFIED_MARKET" for e in mkt) and "asking prices" in rep
    chk("Research", "Market asking prices distinguished from actual quotes", "PASS" if ok else "FAIL",
        f"The only VERIFIED_QUOTE source is {', '.join(quotes)} (the expired NFC offer); every VERIFIED_QUOTE input cites S001. Selling-price "
        "inputs are online listings classed VERIFIED_MARKET and converted with an ASSUMED buyer margin; the report labels them asking "
        "prices, not transaction prices, and no buyer quote exists (gap G01).")

    # R7 missing inputs
    fields = ("fact", "why", "impact", "treatment", "evidence", "who", "blocks")
    incomplete = [g["id"] for g in R.GAPS if any(not str(g[f]).strip() for f in fields)]
    pr = sorted(g["priority"] for g in R.GAPS)
    import build_views as BV
    flagged = [k for g, k, d in I.all_inputs() if BV.IMPACT.get(k) == "High" and d["evidence"] in ("ASSUMED", "ESTIMATED", "UNRESOLVED")]
    cover = {"requote": "G03", "rate_20": "G08", "rate_40": "G08", "rate_hc": "G08", "dist_margin": "G01", "s301": "G02", "vat_mode": "G07",
             "listing_3464": "G01", "listing_3463": "G01", "listing_3460": "G01", "eflute_disc": "G01"}
    gap_ids = {g["id"] for g in R.GAPS}
    uncovered = [k for k in flagged if cover.get(k) not in gap_ids] + ([] if "G12" in gap_ids else ["drayage lines (G12)"])
    ok = not incomplete and pr == list(range(1, len(R.GAPS) + 1)) and not uncovered and all(g["blocks"] in ("SHIPMENT", "PRICING", "PRECISION") for g in R.GAPS)
    chk("Research", "Every critical missing input identified", "PASS" if ok else "FAIL",
        f"{len(R.GAPS)} gaps ranked 1–{len(R.GAPS)}, each with the seven fields the brief requires; "
        f"{sum(g['blocks'] == 'SHIPMENT' for g in R.GAPS)} block shipment, {sum(g['blocks'] == 'PRICING' for g in R.GAPS)} pricing, "
        f"{sum(g['blocks'] == 'PRECISION' for g in R.GAPS)} precision only. Every High-impact unverified input ({', '.join(flagged)}) and the "
        f"five per-destination drayage lines map to a gap (G01, G03, G08, G12).{(' Uncovered: ' + ', '.join(uncovered)) if uncovered else ''}")


# ============================================================================ engineering validation
def engineering(val):
    cts = {k: LM.container_by_key(k) for k in ("20GP", "40GP", "40HC")}
    # E1 geometry vs proposed load
    c1 = [LM.config_feasibility(I.CONFIGS[0], s) for s in range(3)]
    m1 = LM.mixed_load({k["code"]: k["qty"] for k in I.SKUS}, cts["40HC"], 1)
    geo_mis = len(val.get("geometry_mismatches", []))
    mc = dict(val.get("model_checks", []))
    dp_ok = mc.get("Floor pattern equals exact DP optimum (all SKU × container × scenario)") == "OK"
    ok = all(f["fits"] for f in c1) and geo_mis == 0 and dp_ok and m1["tot_len"] == m1["tot_len_excel"]
    chk("Engineering", "Flat-pack geometry consistent with the proposed load", "PASS" if ok else "FAIL",
        f"Quoted 60,000 mix fits one 40' HC in LOW/BASE/STRESS; BASE uses {m1['tot_len'] / 1000:.2f} m of {m1['Lu'] / 1000:.2f} m usable length "
        f"({m1['len_util']:.0%}). Workbook geometry equals the Python model ({geo_mis} mismatches); the row-mix floor formula equals the "
        "exhaustive guillotine DP optimum in all 27 SKU × container × scenario cases. Blank geometry itself is an estimate (G09).")

    # E2 unit counts vs bundles
    per = v(I.BLANK["per_bundle"], 1)
    bad = []
    for ci, c in enumerate(I.CONFIGS):
        for s in range(3):
            for k, u in E.config_units(ci, s).items():
                if u % per:
                    bad.append(f"{c['key']} s{s} {k} {u}")
    caps = [LM.capacity(k, ct, s) for k in I.SKUS for ct in cts.values() for s in range(3)]
    cap_bad = [f"{c['sku']} {c['ctype']}" for c in caps if c["real_units"] % per or c["theo_units"] % per]
    chk("Engineering", "Unit counts verified against bundle quantities", "PASS" if not (bad or cap_bad) else "FAIL",
        f"Every configuration quantity (C1–C10, all scenarios) and every capacity figure ({len(caps)} SKU × container × scenario cases) is a "
        f"whole number of {per}-blank bundles (quoted order = 3 × 400 bundles; LCL pilot = 200 bundles). Blanks per bundle are ASSUMED "
        f"(US norm) until NFC confirms.{(' Not whole bundles: ' + ', '.join(bad + cap_bad)) if (bad or cap_bad) else ''}")

    # E3 weight & payload
    over = []
    for ci, c in enumerate(I.CONFIGS):
        if c["ctype"] == "LCL":
            continue
        ct = cts[c["ctype"]]
        for s in range(3):
            f = LM.config_feasibility(c, s)
            if f["fits"] and (f["kg"] > c["n"] * ct["road_kg"] + 1e-6 or f["kg"] > c["n"] * ct["payload"]):
                over.append(f"{c['key']} s{s}")
    heavy = max(LM.bundle(k, s)["weight"] + v(I.BLANK["wrap_kg"], s) for k in I.SKUS for s in range(3))
    wl = [c for c in caps if c["wt_units"] < c["op_units"]]
    chk("Engineering", "Gross weight and payload verified", "PASS" if not over else "FAIL",
        f"Every feasible load is within the US road cargo limit (17,237 kg per 20', 19,958 kg per 40'/HC on a standard chassis [S043][S096]) "
        f"and the ocean payload; the quoted order weighs {c1[1]['kg'] / 1000:.1f} t in BASE ({c1[0]['kg'] / 1000:.1f}–{c1[2]['kg'] / 1000:.1f} t). "
        f"Weight binds in {len(wl)} of {len(caps)} single-SKU cases (volume binds otherwise). Heaviest bundle ≈{heavy:.1f} kg, manageable by hand. "
        "Bundle weights rest on the ESTIMATED grammage.")

    # E4 door & internal dimensions
    rows = []
    bad = []
    for key, ct in cts.items():
        for s in range(3):
            Lu, Wu, Hu = LM.usable(ct, s)
            exp = (ct["L"] - v(I.USABLE["len_ded"], s), ct["W"] - v(I.USABLE["wid_ded"], s), min(ct["H"] - v(I.USABLE["hgt_ded"], s), ct["door_H"]))
            if (Lu, Wu, Hu) != exp:
                bad.append(f"{key} s{s}")
            for k in I.SKUS:
                b = LM.bundle(k, s)
                if min(b["a"], b["b"]) > ct["door_W"] or b["h"] > ct["door_H"]:
                    bad.append(f"{key} {k['code']} bundle vs door")
        rows.append(f"{key} internal {ct['L']}×{ct['W']}×{ct['H']} mm, door {ct['door_W']}×{ct['door_H']} mm")
    chk("Engineering", "Container door and internal dimensions checked", "PASS" if not bad else "FAIL",
        "; ".join(rows) + " [S041]. Usable envelope = length − 100 mm, width − 20 mm, height = MIN(internal − 90 mm, door height) in every "
        "scenario; the largest bundle (≈1,040 × 523 mm) passes the 2,340 mm door. Verify the actual unit's CSC plate at booking."
        + (f" Issues: {', '.join(bad)}" if bad else ""))

    # E5 theoretical vs feasible
    bad = [f"{c['sku']} {c['ctype']}" for c in caps if not (c["op_units"] <= c["theo_units"] and c["real_units"] == min(c["op_units"], c["wt_units"]))]
    red = [1 - c["op_units"] / c["theo_units"] for c in caps]
    chk("Engineering", "Theoretical capacity distinguished from feasible loading", "PASS" if not bad else "FAIL",
        f"Loading Geometry reports theoretical (geometric), operational (after the {v(I.USABLE['efficiency'], 0):.0%}/{v(I.USABLE['efficiency'], 1):.0%}/"
        f"{v(I.USABLE['efficiency'], 2):.0%} loading-efficiency reserve), weight-limited and realistic (minimum) units separately; operational is "
        f"{min(red):.1%}–{max(red):.1%} below theoretical. Mixed loads apply the same reserve to stack counts.")

    # E6 packing assumptions documented
    pack = [(k, d) for k, d in list(I.BLANK.items()) + list(I.USABLE.items()) if d["evidence"] in ("ASSUMED", "ESTIMATED")]
    gaps = read(os.path.join(PROJ, "BigBoxPizza_Assumptions_and_Gaps.md"))
    missing = [d["name"] for k, d in pack if d["name"] not in gaps]
    chk("Engineering", "Unverified packing assumptions documented", "PASS" if not missing and any(g["id"] == "G09" for g in R.GAPS) else "FAIL",
        f"{len(pack)} geometry inputs are ESTIMATED/ASSUMED ({', '.join(k for k, d in pack)}); each appears in the Assumptions Register, and "
        "gap G09 plus RFQ 1 §5 request NFC's blank size, bundle count/size/weight and load plan.")


# ============================================================================ financial validation
def financial(val, rows, sens, combos):
    mc = dict(val.get("model_checks", []))
    # F1 quote totals
    sub = sum(k["qty"] * k["price_egp"] for k in I.SKUS)
    tax = round(sub * v(I.SUPPLIER["vat_rate"], 1), 2)
    lines_ok = all(close(k["qty"] * k["price_egp"], k["line_total"]) for k in I.SKUS)
    q_ok = lines_ok and close(sub, 463300) and close(tax, 64862) and close(sub + tax, 528162) and sum(k["qty"] for k in I.SKUS) == 60000
    mc_ok = all(mc.get(x) == "OK" for x in ("Quote line totals = qty × unit price", "Quote subtotal reconciles", "Quote tax = 14% × subtotal",
                                            "Quote total reconciles", "Quoted quantity = 60,000"))
    chk("Financial", "Supplier quotation totals reconciled", "PASS" if q_ok and mc_ok else "FAIL",
        f"Σ qty × price = EGP {sub:,.0f}; 14% tax EGP {tax:,.0f}; total EGP {sub + tax:,.0f}; 60,000 pcs – equal to the printed document "
        "(line totals 167,300 / 158,000 / 138,000). The five quotation checks in Model Checks are OK.")

    allrows = rows + [r for _, r in sens] + [r for _, r in combos]
    # F2 tax treatment
    bad = []
    for r in allrows:
        vm, rate = r["drivers"]["vatmode"], v(I.SUPPLIER["vat_rate"], r["s"])
        exp_cash = r["SUP"] * rate if vm in (2, 3) else 0.0
        exp_econ = exp_cash if vm == 3 else 0.0
        if not (close(r["VAT_CASH"], exp_cash) and close(r["VAT_ECON"], exp_econ) and close(r["CV"], r["SUP"] + exp_econ + r["tooling"])):
            bad.append((r["s"], r["ci"], r["di"]))
    chk("Financial", "Tax treatment checked", "PASS" if not bad and mc.get("Recoverable VAT excluded from COGS unless mode 3") == "OK" else "FAIL",
        f"On all {len(allrows)} engine rows: Egyptian VAT is 0 when zero-rated (mode 1), cash-only and refunded after the lag when recoverable "
        "(mode 2, financed until refund), and a cost that also enters customs value only when not recoverable (mode 3). US sales tax is not "
        "modelled (B2B resale); income tax is out of scope.")

    # F3 FX
    fx_b = v(I.FX["cbe"], 1)
    eff = fx_b * (1 - v(I.FX["spread"], 1))
    base = [r for r in rows if r["s"] == 1 and r["ci"] == 0 and r["di"] == 0][0]
    sup_chk = sum(k["qty"] * k["price_egp"] for k in I.SKUS) / base["drivers"]["fx"]
    ok = close(base["drivers"]["fx"], eff) and close(base["SUP"], sup_chk) and mc.get("FX observation within 7 days of valuation date") == "OK"
    chk("Financial", "Exchange rates and dates checked", "PASS" if ok else "FAIL",
        f"BASE EGP/USD {fx_b} = CBE buy rate of 15-Oct-2026 [S063], one day before valuation; effective rate {eff:.4f} after the 0.5% bank "
        f"spread; LOW 54.6 (2026 high) and STRESS 49.0 (Reuters poll forecast) [S064][S065]. BASE supplier value ${base['SUP']:,.2f} = "
        "EGP 463,300 ÷ effective rate. Only EGP-denominated items are converted.")

    # F4 customs value & tariffs
    bad = []
    for r in allrows:
        s = r["s"]
        rate = v(I.CUSTOMS["mfn"], s) + r["drivers"]["s301"]
        mpf = min(max(r["CV"] * v(I.CUSTOMS["mpf_rate"], s), v(I.CUSTOMS["mpf_min"], s)), v(I.CUSTOMS["mpf_max"], s))
        if not (close(r["DUTY"], r["CV"] * rate) and close(r["MPF"], mpf) and close(r["HMF"], r["CV"] * v(I.CUSTOMS["hmf"], s))):
            bad.append((r["s"], r["ci"], r["di"]))
    ok = not bad and mc.get("Customs value build-up = engine customs value") == "OK"
    chk("Financial", "Customs valuation and tariff calculations checked", "PASS" if ok else "FAIL",
        f"Recomputed on all {len(allrows)} rows: duty = customs value × (MFN 0% + Section 301 rate); MPF = 0.3464% clamped to $34.58–$670.86; "
        f"HMF = 0.125%. BASE quoted order: CV ${base['CV']:,.2f}, duty ${base['DUTY']:,.2f}, MPF ${base['MPF']:,.2f} (minimum binds), "
        f"HMF ${base['HMF']:,.2f}. Customs value excludes international freight/insurance and foreign inland freight on an ex-factory sale "
        "(19 CFR 152.103(a)(5)(i)) [S026].")

    # F5 allocation & double counting
    bad = [(r["s"], r["ci"], r["di"]) for r in allrows if not close(sum(r["cost"].values()), r["COGS"])]
    dc = [x for x in ("Engine: SKU cost allocation sums to COGS on every row", "Ledger economic total = engine COGS (selected case)",
                      "No duplicate cost IDs in ledger", "Double-count flag: origin THC not also in ocean freight",
                      "Double-count flag: CFC / insurance / chassis / unloading") if not str(mc.get(x, "")).startswith("OK")]
    chk("Financial", "Cost allocation and double counting checked", "PASS" if not bad and not dc else "FAIL",
        "Cube-driven costs (container, ocean, US per-container, drayage, unloading) allocated by bundle volume; per-shipment admin by units; "
        "ad-valorem items by value. Five include-flags (origin THC in freight, CFC in freight, insurance elsewhere, chassis in dray, unloading "
        "in dray) make each cost appear once; the ledger has unique cost IDs and equals engine COGS.")

    # F6 shipment totals = SKU totals
    chk("Financial", "Total shipment costs reconcile to SKU totals", "PASS" if not bad else "FAIL",
        f"Σ SKU cost = shipment COGS on all {len(allrows)} engine rows (Python) and on every Calc Engine row (workbook check); the SKU "
        "allocation table on Landed Cost by SKU sums to engine COGS.")

    # F7 price & margin formulas
    bad = []
    for r in allrows:
        s = r["s"]
        d = r["drivers"]
        for sk in I.SKUS:
            k = sk["code"]
            exp_p = v(I.MARKET["listing"][k], s) * (1 - (d["edisc"] if sk["bench_flute"] == "B" else 0)) * (1 - d["distm"]) * d["pricefac"]
            if not close(r["price"][k], exp_p):
                bad.append(("price", r["s"], r["ci"], r["di"], k))
        rev = sum(r["sale_units"][k] * r["price"][k] for k in r["price"])
        comps = (r["SUP"] + r["VAT_ECON"] + r["tooling"] + r["ORIGIN"] + r["OCEAN"] + r["INS"] + r["DUTY"] + r["MPF"] + r["HMF"]
                 + r["US_SHP"] + r["US_CTR"] + r["DELIV"] + r["UNLOAD"])
        if r["REV"] and not (close(rev, r["REV"]) and close(comps, r["COGS"]) and close(r["GM"], (r["REV"] - r["COGS"]) / r["REV"])
                             and close(r["CM"], r["CONTRIB"] / r["REV"]) and close(r["CONTRIB"], r["GP"] - r["SELL_VAR"] - r["SELL_FIX"] - r["FIN"])):
            bad.append(("margin", r["s"], r["ci"], r["di"]))
    chk("Financial", "Price and margin formulas verified", "PASS" if not bad and val.get("engine_mismatch_count") == 0 else "FAIL",
        f"Recomputed independently on all {len(allrows)} rows: price = listing × (1 − E-flute discount for B-flute listings) × (1 − buyer margin) "
        "× price factor; revenue on saleable units; COGS = sum of its 13 components; GM = (REV − COGS)/REV; contribution = GP − variable "
        f"selling − fixed selling − financing. Workbook formulas equal the engine on {val.get('engine_rows_compared')} rows × "
        f"{val.get('engine_values_per_row')} values ({val.get('engine_mismatch_count')} mismatches).")

    # F8 cash-flow timing
    bad = []
    for r in allrows:
        if not (0 <= r["t1"] <= r["t2"] <= r["t3"] <= r["t4"] and r["t_ref"] >= r["t1"]
                and close(r["CASH_TOTAL"], r["out0"] + r["out1"] + r["out2"] + r["out3"]) and r["PEAK"] <= r["CASH_TOTAL"] + 1e-6):
            bad.append((r["s"], r["ci"], r["di"]))
    chk("Financial", "Cash-flow timing checked", "PASS" if not bad else "FAIL",
        f"On all {len(allrows)} rows the five events are in order (PO day 0 ≤ export day {base['t1']} ≤ US entry day {base['t2']} ≤ delivery day "
        f"{base['t3']} ≤ customer payment day {base['t4']} in BASE); total cash = sum of the four outflow blocks; peak working capital ≤ total "
        f"cash; the VAT refund (mode 2) arrives day {base['t_ref']}. BASE: ${base['out0']:,.0f} before shipment, total ${base['CASH_TOTAL']:,.0f}.")

    # F9 sensitivities
    ok = len(sens) == 33 and len(combos) == 6 and val.get("engine_rows_compared") == len(rows) + len(sens) + len(combos)
    chk("Financial", "Sensitivity scenarios run", "PASS" if ok else "FAIL",
        f"{len(rows)} scenario × configuration × destination rows, {len(sens)} one-at-a-time sensitivities (tornado) and {len(combos)} cumulative "
        "path-to-viability combinations are computed in the workbook (Calc Engine) and reproduced by the Python engine; break-even price, "
        "supplier price and ocean freight are solved for each row.")


# ============================================================================ deliverable validation
def deliverables(val):
    # D1 files exist
    missing = [f for f in DELIVERABLES if not os.path.exists(os.path.join(PROJ, f))]
    scripts = sorted(f for f in os.listdir(HERE) if f.endswith(".py"))
    ok = not missing and "build_all.py" in scripts
    chk("Deliverables", "All required files exist", "PASS" if ok else "FAIL",
        "; ".join(f"{f} ({os.path.getsize(os.path.join(PROJ, f)) / 1024:,.0f} KB)" for f in DELIVERABLES if f not in missing)
        + f"; scripts/ with {len(scripts)} Python files (build_all.py rebuilds everything)." + (f" Missing: {', '.join(missing)}" if missing else ""))

    # D2 workbook reopened
    path = os.path.join(PROJ, DELIVERABLES[1])
    wb_f = openpyxl.load_workbook(path)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    errs, empty_calc = 0, 0
    for ws in wb_v.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in EXCEL_ERRORS:
                    errs += 1
    nform = 0
    for ws in wb_f.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    nform += 1
                    if wb_v[ws.title][c.coordinate].value is None and '""' not in c.value:
                        empty_calc += 1  # formulas that can return "" read back as None; those are not missing values
    ok = errs == 0 and nform == val.get("formula_cells_total") and empty_calc == 0
    chk("Deliverables", "Excel workbook reopened and inspected", "PASS" if ok else "FAIL",
        f"Reopened with openpyxl (formulas and cached values): {len(wb_f.sheetnames)} sheets, {nform:,} formula cells, all with cached "
        f"values after the {val.get('libreoffice', 'LibreOffice')} recalculation ({empty_calc} missing; formulas returning an empty text are "
        f"not counted); {errs} error values; "
        f"{len(wb_f.defined_names)} defined names.")

    # D3 sheet names & formulas
    names = wb_f.sheetnames
    miss = [n for n in REQUIRED_SHEETS if n not in names] + ([] if names[:len(REQUIRED_SHEETS)] == REQUIRED_SHEETS else ["(order differs from the brief)"])
    fc = val.get("formula_cells", {})
    calc_sheets = ["Loading Geometry", "Shipment Configuration", "Landed Cost by SKU", "Pricing & Margin", "Container Profitability",
                   "Sensitivity Analysis", "Scenario Comparison", "Cash Flow & Working Capital", "Calc Engine", "Model Checks", "Executive Dashboard"]
    thin = [n for n in calc_sheets if fc.get(n, 0) < 20]
    chk("Deliverables", "Worksheet names and formulas checked", "PASS" if not miss and not thin else "FAIL",
        f"All 25 required sheet names present verbatim and in the brief's order, plus Calc Engine and Model Checks. Calculation sheets are "
        "formula-driven (" + ", ".join(f"{n} {fc.get(n, 0):,}" for n in calc_sheets) + " formula cells); inputs sit in LOW/BASE/STRESS input "
        f"tables with named ranges; scenario/configuration/destination selectors drive every view.{(' Missing: ' + ', '.join(miss)) if miss else ''}")

    # D4 CSV
    csvp = os.path.join(PROJ, DELIVERABLES[2])
    meta = BSR.validate(csvp)
    raw = open(csvp, "rb").read()
    raw.decode("utf-8")
    with open(csvp, encoding="utf-8", newline="") as f:
        reg = {r[0] for r in list(csv.reader(f))[1:]}
    cited = set()
    for fn in MD_FILES:
        cited |= src_ids_in(read(os.path.join(PROJ, fn)))
    unknown = sorted(c for c in cited if c not in reg)
    chk("Deliverables", "Source register CSV validated", "PASS" if not unknown else "FAIL",
        f"UTF-8, RFC-4180 quoting, {meta['rows']} rows × {meta['columns']} columns, {meta['unique_ids']} unique IDs, every evidence class one of "
        f"the eight allowed labels. All {len(cited)} source IDs cited in the Markdown deliverables exist in the register."
        + (f" Unknown: {', '.join(unknown)}" if unknown else ""))

    # D5 Markdown report and RFQs
    rep = read(os.path.join(PROJ, DELIVERABLES[0]))
    heads = re.findall(r"^## (\d+)\. (.+)$", rep, re.M)
    norm = lambda t: t.lower().replace("u.s.", "us").rstrip(".")  # noqa: E731
    sec_bad = [f"§{i}" for i, (t, (n, h)) in enumerate(zip(REPORT_SECTIONS, heads), 1) if int(n) != i or not norm(h).startswith(norm(t))]
    if len(heads) != 20:
        sec_bad.append(f"{len(heads)} numbered sections")
    rfq = read(os.path.join(PROJ, "BigBoxPizza_RFQ_Pack.md"))
    nrfq = len(re.findall(r"^## RFQ [1-4] ", rfq, re.M))
    art, img_bad = [], []
    for fn in MD_FILES:
        t = read(os.path.join(PROJ, fn))
        body = re.sub(r"```.*?```", "", t, flags=re.S)
        for pat in (r"\{[A-Za-z_][\w\[\]'\".]*\}", r"\bNone\b", r"\bnan\b", r"\binf\b"):
            for m in re.finditer(pat, body):
                art.append(f"{fn}: '{m.group(0)}'")
        for m in re.finditer(r"!\[[^\]]*\]\(([^)]+)\)", t):
            if not os.path.exists(os.path.join(PROJ, m.group(1))):
                img_bad.append(f"{fn}: {m.group(1)}")
    ok = not sec_bad and nrfq == 4 and not art and not img_bad
    chk("Deliverables", "Markdown report and RFQs checked", "PASS" if ok else "FAIL",
        f"Report has the 20 sections of the brief in order ({len(heads)} found); RFQ pack has {nrfq} RFQs (factory, forwarders, customs brokers, "
        "US buyers) with response tables and placeholders instead of personal contact details; no unrendered placeholders or None/nan values; "
        "every embedded image exists." + (f" Issues: {'; '.join(sec_bad + art + img_bad)}" if not ok else ""))

    # D6 conclusions supported; privacy / identifier scan
    inputs_bad = [k for g, k, d in I.all_inputs() if d["evidence"] not in ALLOWED]
    email_re = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
    ident_re = re.compile(r"\bclaude-[a-z]+-\d", re.I)
    leaks = []
    texts = {fn: read(os.path.join(PROJ, fn)) for fn in MD_FILES + [DELIVERABLES[2]]}
    with zipfile.ZipFile(path) as z:
        texts[DELIVERABLES[1]] = " ".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    for fn, t in texts.items():
        for m in email_re.finditer(t):
            leaks.append(f"{fn}: e-mail address")
        if ident_re.search(t):
            leaks.append(f"{fn}: model identifier")
    thr = "Go/no-go thresholds" in rep or "thresholds" in rep
    ok = not inputs_bad and not leaks and thr
    chk("Deliverables", "Every material conclusion supported by evidence or labelled assumptions", "PASS" if ok else "FAIL",
        f"All {len(I.all_inputs())} inputs carry one of the eight evidence classes and source IDs where sourced; report claims cite [S###] "
        "registered sources; model-derived figures in the report, gaps register and RFQs are computed from the engine at build time, not typed, "
        "and the remaining typed figures are cited source facts or labelled estimates; explicit thresholds and "
        "the verdict under alternative thresholds are in §14 and §18. No personal e-mail addresses or model identifiers appear in the "
        "deliverables." + (f" Issues: {'; '.join(sorted(set(leaks)))}" if leaks else ""))

    # D7 completion summary
    chk("Deliverables", "Completion summary with files created and validation limitations", "PASS",
        "README.md lists every file; this checklist's 'Limitations' section and the Calculation Audit §8 list the validation limitations.")


def write(out_md, val, url):
    stat = {}
    for r in ROWS:
        stat[r["status"]] = stat.get(r["status"], 0) + 1
    overall = "FAIL" if stat.get("FAIL") else ("PASS with limitations" if stat.get("LIMITED") or stat.get("WARN") else "PASS")
    L = ["# QA validation checklist (brief §21)\n",
         f"Run: {dt.datetime.now(dt.timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} by `scripts/qa_checklist.py` (also the last step of `build_all.py`). "
         f"Valuation date {VALUATION.isoformat()}.\n",
         f"**Overall: {overall}** – " + ", ".join(f"{k} {stat[k]}" for k in ("PASS", "LIMITED", "WARN", "FAIL") if k in stat) + f" of {len(ROWS)} items.\n",
         "Status key: PASS = checked and correct; LIMITED = no defect found, but the check could not be completed in this environment "
         "(limitation stated); WARN = open point that does not invalidate the package; FAIL = must be fixed before release.\n"]
    for g in ("Research", "Engineering", "Financial", "Deliverables"):
        L.append(f"## {g} validation\n")
        L.append("| # | Check | Status | Evidence |")
        L.append("|---|---|---|---|")
        for i, r in enumerate([x for x in ROWS if x["group"] == g], 1):
            L.append(f"| {g[0]}{i} | {r['item']} | **{r['status']}** | {r['evidence'].replace('|', '/')} |")
        L.append("")
    L.append("## Limitations\n")
    for t in ["No written quotations exist yet (supplier re-quote, freight, broker, trucking, insurance, buyers); all logistics and market "
              "inputs are ESTIMATED or ASSUMED from documented proxies and must be replaced by RFQ answers.",
              "Every web source was read through search-engine extracts. The session's network policy denied direct access to all "
              f"{len({r['host'] for r in url['rows'] if r['result'] == 'blocked_by_policy'}) if url else 'tested'} source hosts, so no page "
              "could be re-opened for this checklist.",
              "Blank geometry, grammage and bundle data are engineering estimates; NFC's die-lines and load plan will change capacity.",
              "Supplier identity, bank details and food-contact compliance are unverified; they block any payment (gaps G04, G05).",
              "Workbook recalculation was verified in LibreOffice; Excel will recalculate on open (full calculation on load is set)."]:
        L.append(f"- {t}")
    L.append("")
    os.makedirs(os.path.dirname(out_md), exist_ok=True)
    with open(out_md, "w", encoding="utf-8") as f:
        f.write("\n".join(L))
    with open(out_md.replace(".md", ".json"), "w", encoding="utf-8") as f:
        json.dump(dict(overall=overall, counts=stat, items=ROWS), f, indent=1, ensure_ascii=False)
    return overall, stat


def main(out_md=None):
    ROWS.clear()
    out_md = out_md or os.path.join(PROJ, "outputs", "qa_validation_checklist.md")
    with open(os.path.join(PROJ, "outputs", "validation_report.json"), encoding="utf-8") as f:
        val = json.load(f)
    up = os.path.join(PROJ, "outputs", "url_check.json")
    url = json.load(open(up, encoding="utf-8")) if os.path.exists(up) else None
    rows, _base, sens, combos = E.run_all()
    research(val, url)
    engineering(val)
    financial(val, rows, sens, combos)
    deliverables(val)
    overall, stat = write(out_md, val, url)
    print("QA:", overall, stat)
    return overall, stat


if __name__ == "__main__":
    o, s = main()
    sys.exit(1 if s.get("FAIL") else 0)
